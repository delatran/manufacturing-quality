"""Case knowledge checks; local tests use synthetic in-memory HTTP payloads only."""

from __future__ import annotations

import hashlib
import json
import os
from decimal import Decimal
from pathlib import Path

import httpx
import pytest
from analyst.case_study.analytics import (
    Scope,
    analysis_report,
    filter_events,
    filter_products,
    kpis,
)
from analyst.case_study.ingestion import SourceError, fetch_events
from analyst.case_study.pipeline import execute_case_study
from analyst.case_study.preparation import build_snapshot, parse_cost, parse_source_date

MASTER = b"product_id,product_name,category,production_line,monthly_output\n1,Widget A,Category 1,Line-1,100\n2,Widget B,Category 1,Line-1,200\n"


def event(identifier=1, **updates):
    return {
        "defect_id": identifier,
        "product_id": 1,
        "defect_type": "Functional",
        "defect_date": "2024-01-15",
        "defect_location": "Component",
        "severity": "Critical",
        "inspection_method": "Visual Inspection",
        "repair_cost": 10,
        **updates,
    }


def snapshot(rows):
    return build_snapshot(
        rows,
        {
            "source_sha256": hashlib.sha256(json.dumps(rows).encode()).hexdigest(),
            "source_url": "https://synthetic.invalid/events",
            "fetched_at": "2026-01-01T00:00:00Z",
        },
        MASTER,
    )


def test_unknown_cost_and_unmatched_counts_preserve_grain():
    data = snapshot([event(), event(2, repair_cost=None), event(3, product_id=99, repair_cost=-3)])
    result = analysis_report(data, Scope())["kpis"]
    assert result["event_count"] == 3
    assert result["known_repair_cost"] == 10
    assert result["cost_unknown_count"] == 2
    assert result["linked_event_count"] == 2
    assert result["reference_output"] == 1800  # Both master products, six months.
    assert result["reference_intensity_per_1000"] == pytest.approx(2 / 1800 * 1000)


def test_distinct_observed_product_ids_and_unmatched_events_have_different_populations():
    data = snapshot(
        [
            event(1),
            event(2, product_id=99),
            event(3, product_id=99),
            event(4, product_id=99),
            event(5, product_id=88),
            event(6, product_id=88),
        ]
    )
    report = analysis_report(data, Scope())
    values = report["kpis"]
    eligible_ids_with_events = {item.product_id for item in data.events if item.matched_product}
    assert values["products_with_events"] == 3
    assert values["cohort_product_count"] == 2
    assert values["linked_event_count"] == 1 and values["unmatched_event_count"] == 5
    assert len(eligible_ids_with_events) == 1
    assert values["products_with_events"] - values["unmatched_event_count"] != len(
        eligible_ids_with_events
    )
    assert "including identifiers unmatched" in report["definitions"]["products_with_events"]
    assert "no selected events" in report["definitions"]["cohort_product_count"]
    assert "not distinct products" in report["definitions"]["linked_event_count"]
    assert "cannot be subtracted" in report["definitions"]["unmatched_event_count"]
    empty = analysis_report(data, Scope(start_month="2024-03", end_month="2024-03"))
    assert empty["kpis"]["products_with_events"] == 0 and empty["kpis"]["cohort_product_count"] == 2
    assert empty["definitions"] == report["definitions"]


def test_calendar_evidence_uses_actual_filtered_dates_and_empty_months():
    data = snapshot(
        [
            event(1, defect_date="2024-01-28"),
            event(2, defect_date="2024-01-31", severity="Minor"),
            event(3, defect_date="2024-02-15"),
            event(4, defect_date="2024-02-29", severity="Minor"),
        ]
    )
    whole = analysis_report(data, Scope(start_month="2024-01", end_month="2024-02"))
    scoped = analysis_report(
        data, Scope(start_month="2024-01", end_month="2024-02", severity="Critical")
    )
    coverage = scoped["calendar_coverage"]
    assert coverage["scope"] == scoped["scope"] and coverage["scope_kind"] == "selected_scope"
    assert (
        coverage["calendar_days"] == 60
        and coverage["days_with_events"] == 2
        and coverage["days_without_events"] == 58
    )
    assert coverage["max_observed_day_of_month"] == 28
    assert [row["last_observed_date"] for row in coverage["monthly"]] == [
        "2024-01-28",
        "2024-02-15",
    ]
    assert [row["last_observed_day"] for row in coverage["monthly"]] == [28, 15]
    assert [row["calendar_days_after_last_observed"] for row in coverage["monthly"]] == [3, 14]
    assert whole["calendar_coverage"]["max_observed_day_of_month"] == 31
    assert "after day 28" in scoped["metadata"]["limitations"][-1]
    assert "after day 28" not in whole["metadata"]["limitations"][-1]
    assert scoped["quality"]["temporal_coverage"]["days_with_events"] == 4
    empty = analysis_report(
        data, Scope(start_month="2024-03", end_month="2024-03", severity="Critical")
    )
    assert empty["calendar_coverage"]["max_observed_day_of_month"] is None
    assert empty["calendar_coverage"]["monthly"][0]["last_observed_date"] is None
    assert empty["calendar_coverage"]["days_without_events"] == 31
    assert "after day" not in empty["metadata"]["limitations"][-1]


def test_conflict_quarantines_all_variants_and_complementary_cost_is_observed():
    rows = [
        event(),
        event(),
        event(2, product_id=1),
        event(2, product_id=2),
        event(3, repair_cost=None),
        event(3, repair_cost=12),
    ]
    data = snapshot(rows)
    assert [e.defect_id for e in data.events] == [1, 3]
    assert data.quality["reconciliation"]["raw_rows"] == 6
    assert data.quality["reconciliation"]["duplicate_copies_removed"] == 2
    assert data.quality["reconciliation"]["quarantined_rows"] == 2
    assert data.quality["known_repair_cost"] == 22


def test_invalid_variant_cannot_silently_resolve_same_id():
    data = snapshot([event(), event(defect_date="impossible")])
    assert not data.events
    assert data.quality["quarantined_count"] == 2


def test_date_rules_are_explicit_and_no_year_guess():
    assert parse_source_date("6/3/2024")[0].isoformat() == "2024-06-03"
    assert parse_source_date("6/3/2024")[1] == "ambiguous_date_assumption"
    assert parse_source_date("25/6/2024")[0].isoformat() == "2024-06-25"
    data = snapshot([event(defect_date="2026-06-03")])
    assert data.quality["out_of_scope_count"] == 1
    assert data.events[0].defect_date.year == 2026


@pytest.mark.parametrize(
    "value,status",
    [
        (None, "missing"),
        (-1, "negative"),
        (0, "known"),
        ("1,000", "invalid"),
        (True, "invalid"),
        (float("nan"), "invalid"),
        (10**400, "invalid"),
        (-(10**400), "negative"),
    ],
)
def test_cost_semantics(value, status):
    assert parse_cost(value)[1] == status


def test_oversized_costs_preserve_countable_events_as_unknown():
    data = snapshot([event(1, repair_cost=10**400), event(2, repair_cost=-(10**400))])
    assert len(data.events) == 2
    assert all(item.repair_cost is None for item in data.events)
    assert data.quality["known_repair_cost"] == 0


def test_http_source_stays_in_memory_and_hashes_response(tmp_path, monkeypatch):
    payload = json.dumps([event()]).encode()
    transport = httpx.MockTransport(lambda request: httpx.Response(200, content=payload))
    monkeypatch.chdir(tmp_path)
    rows, provenance = fetch_events("https://synthetic.invalid/events", transport=transport)
    assert rows == [event()]
    assert provenance["source_sha256"] == hashlib.sha256(payload).hexdigest()
    assert list(tmp_path.iterdir()) == []


def test_bad_http_schema_fails_without_local_fallback():
    transport = httpx.MockTransport(lambda request: httpx.Response(200, json=[{"defect_id": 1}]))
    with pytest.raises(SourceError, match="missing fields"):
        fetch_events("https://synthetic.invalid/events", transport=transport)


def test_event_attributes_do_not_shrink_reference_cohort():
    data = snapshot([event(), event(2, severity="Minor")])
    scope = Scope(severity="Critical")
    result = kpis(filter_events(data, scope), filter_products(data, scope), len(scope.months()))
    assert result["event_count"] == 1
    assert result["cohort_product_count"] == 2
    assert result["reference_output"] == 1800


@pytest.fixture
def workbook_paths(tmp_path):
    template_bytes = (
        Path(__file__).resolve().parents[1] / "DATA_QUALITY_TEMPLATE.xlsx"
    ).read_bytes()
    template = tmp_path / "supplied" / "DATA_QUALITY_TEMPLATE.xlsx"
    template.parent.mkdir()
    template.write_bytes(template_bytes)
    output = tmp_path / "artifacts"
    output.mkdir()
    return template, output, template_bytes


def test_workbook_delivery_attests_actual_template_copy_and_finite_readback(workbook_paths):
    from analyst.case_study.workbook import generate_workbook
    from openpyxl import load_workbook

    template, output, template_bytes = workbook_paths
    data = snapshot([event(), event(2, repair_cost=None)])
    generated = generate_workbook(data, output, template)
    delivery = generated["workbook_delivery"]
    completed = output / "DATA_QUALITY.xlsx"
    assert template.read_bytes() == template_bytes
    assert delivery["mode"] == "completed_copy_from_supplied_template"
    assert delivery["template"] == {
        "name": template.name,
        "sha256_before": hashlib.sha256(template_bytes).hexdigest(),
        "sha256_after": hashlib.sha256(template_bytes).hexdigest(),
        "preserved": True,
    }
    assert delivery["output"] == {
        "name": completed.name,
        "sha256": hashlib.sha256(completed.read_bytes()).hexdigest(),
    }
    book = load_workbook(completed, read_only=True, data_only=True)
    try:
        readback = delivery["readback"]
        assert readback["status"] == "verified"
        assert readback["scope"] == "first_four_sheet_names_and_summary_B8"
        assert readback["expected_first_sheet_names"] == [
            "Summary",
            "Issues Found",
            "Assumptions",
            "Impact on Results",
        ]
        assert readback["observed_sheet_names"] == book.sheetnames
        assert book.sheetnames[:4] == readback["expected_first_sheet_names"]
        assert readback["checked_cell"] == "Summary!B8"
        assert (
            readback["expected_analysis_events"]
            == readback["observed_analysis_events"]
            == book["Summary"]["B8"].value
            == 2
        )
        assert (
            "other cells, formatting and visual layout were not certified" in readback["limitation"]
        )
    finally:
        book.close()


@pytest.mark.parametrize("fault", ["summary_count", "sheet_order"])
def test_workbook_delivery_rejects_failed_saved_readback(workbook_paths, monkeypatch, fault):
    from analyst.case_study.workbook import generate_workbook
    from openpyxl.workbook.workbook import Workbook

    template, output, template_bytes = workbook_paths
    save = Workbook.save

    def save_with_fault(book, path):
        if fault == "summary_count":
            book["Summary"]["B8"] = 999
        else:
            book.move_sheet("Summary", offset=1)
        save(book, path)

    monkeypatch.setattr(Workbook, "save", save_with_fault)
    with pytest.raises(AssertionError, match="read-back failed"):
        generate_workbook(snapshot([event()]), output, template)
    assert template.read_bytes() == template_bytes
    assert not (output / "empirical_summary.json").exists()


def test_workbook_delivery_rejects_template_change_during_generation(workbook_paths, monkeypatch):
    from analyst.case_study.workbook import generate_workbook
    from openpyxl.workbook.workbook import Workbook

    template, output, template_bytes = workbook_paths
    save = Workbook.save

    def save_then_change_supplied_copy(book, path):
        save(book, path)
        template.write_bytes(template_bytes + b"changed after template was loaded")

    monkeypatch.setattr(Workbook, "save", save_then_change_supplied_copy)
    with pytest.raises(AssertionError, match="template changed"):
        generate_workbook(snapshot([event()]), output, template)
    assert not (output / "empirical_summary.json").exists()


def test_workbook_generator_refuses_to_overwrite_supplied_template(workbook_paths):
    from analyst.case_study.workbook import generate_workbook

    _, output, template_bytes = workbook_paths
    template = output / "DATA_QUALITY.xlsx"
    template.write_bytes(template_bytes)
    with pytest.raises(ValueError, match="separate from the supplied template"):
        generate_workbook(snapshot([event()]), output, template)
    assert template.read_bytes() == template_bytes


@pytest.mark.parametrize("tamper_output", [False, True])
def test_pipeline_persists_actual_workbook_delivery_and_rejects_changed_output(
    workbook_paths, monkeypatch, tamper_output
):
    from zipfile import ZipFile

    from analyst.case_study import pipeline
    from analyst.case_study.workbook import generate_workbook

    template, output, template_bytes = workbook_paths
    master = template.parent / "master_data.csv"
    master.write_bytes(MASTER)
    rows = [event(), event(2, repair_cost=None)]
    provenance = {
        "source_sha256": hashlib.sha256(json.dumps(rows).encode()).hexdigest(),
        "source_url": "https://synthetic.invalid/events",
        "fetched_at": "2026-01-01T00:00:00Z",
    }
    monkeypatch.setattr(pipeline, "fetch_events", lambda _: (rows, provenance))
    monkeypatch.setattr(pipeline, "EXPECTED_SOURCE_SHA256", provenance["source_sha256"])
    monkeypatch.setattr(pipeline, "EXPECTED_MASTER_SHA256", hashlib.sha256(MASTER).hexdigest())
    # Isolate delivery propagation from the pinned historical-source oracle and report.
    monkeypatch.setattr(pipeline, "evaluate_date_sensitivity", lambda *args: {})
    monkeypatch.setattr(pipeline, "_qa", lambda *args: {"status": "synthetic_delivery_test"})
    monkeypatch.setattr(
        pipeline, "_render_report", lambda _: ("Synthetic workbook delivery test", [])
    )

    def generate_then_maybe_tamper(*args):
        generated = generate_workbook(*args)
        if tamper_output:
            completed = Path(generated["workbook"])
            completed.write_bytes(completed.read_bytes() + b"changed after read-back")
        return generated

    monkeypatch.setattr(pipeline, "generate_workbook", generate_then_maybe_tamper)
    if tamper_output:
        with pytest.raises(AssertionError, match="workbook changed after read-back"):
            pipeline.execute_case_study(output, master, template)
        assert not (output / "evidence.json").exists()
        assert not (output / "evidence.zip").exists()
    else:
        result = pipeline.execute_case_study(output, master, template)
        delivery = result["workbook_delivery"]
        assert (
            json.loads((output / "evidence.json").read_text(encoding="utf-8"))["workbook_delivery"]
            == delivery
        )
        artifact = next(item for item in result["artifacts"] if item["name"] == "DATA_QUALITY.xlsx")
        assert artifact["sha256"] == delivery["output"]["sha256"]
        with ZipFile(output / "evidence.zip") as archive:
            assert json.loads(archive.read("evidence.json"))["workbook_delivery"] == delivery
            assert (
                hashlib.sha256(archive.read("DATA_QUALITY.xlsx")).hexdigest()
                == delivery["output"]["sha256"]
            )
    assert template.read_bytes() == template_bytes
    assert master.read_bytes() == MASTER


@pytest.mark.skipif(
    os.environ.get("RUN_LIVE_CASE_STUDY") != "1",
    reason="Set RUN_LIVE_CASE_STUDY=1 for live HTTP and full artifacts",
)
def test_live_case_full_pipeline_matches_independent_oracle_and_artifacts(tmp_path):
    from zipfile import ZipFile

    from openpyxl import load_workbook

    result = execute_case_study(tmp_path / "artifacts")
    assert result["qa"]["status"] == "verified"
    assert len(result["qa"]["checks"]) == 11
    assert result["analysis"]["kpis"]["event_count"] == 1071
    assert result["analysis"]["kpis"]["critical_count"] == 349
    assert Decimal(str(result["analysis"]["kpis"]["known_repair_cost"])) == Decimal("508864.39")
    assert [row["event_count"] for row in result["analysis"]["monthly"]] == [
        147,
        119,
        148,
        195,
        210,
        252,
    ]
    assert [s["kpis"]["event_count"] for s in result["sensitivity"]["scenarios"].values()] == [
        1071,
        837,
        653,
    ]
    for claim in result["numeric_claims"]:
        assert "[claim:" + claim["id"] + "]" in result["report"]
        assert claim["path"].startswith("$.")
        assert claim["code_sha256"] == result["code_sha256"]
    book = load_workbook(tmp_path / "artifacts/DATA_QUALITY.xlsx", read_only=True)
    assert book.sheetnames[:4] == ["Summary", "Issues Found", "Assumptions", "Impact on Results"]
    assert book["Summary"]["B8"].value == 1071
    book.close()
    delivery = result["workbook_delivery"]
    template_sha256 = hashlib.sha256(
        (Path(__file__).resolve().parents[1] / "DATA_QUALITY_TEMPLATE.xlsx").read_bytes()
    ).hexdigest()
    assert (
        delivery["template"]["sha256_before"]
        == delivery["template"]["sha256_after"]
        == template_sha256
    )
    assert delivery["template"]["preserved"] and delivery["readback"]["status"] == "verified"
    assert delivery["readback"]["observed_analysis_events"] == 1071
    assert (
        delivery["output"]["sha256"]
        == hashlib.sha256((tmp_path / "artifacts/DATA_QUALITY.xlsx").read_bytes()).hexdigest()
    )
    with ZipFile(tmp_path / "artifacts/evidence.zip") as archive:
        assert archive.testzip() is None
        assert {
            "DATA_QUALITY.xlsx",
            "report.md",
            "evidence.json",
            "numeric_claims.json",
            "cleaned_records.csv",
            "monthly_quality.png",
        } <= set(archive.namelist())
    assert not list((tmp_path / "artifacts").glob("*defect*.json"))
    # Use the same independent output guardrail as the SDK runtime, including
    # numeric tokens in the actual exported report rather than only helper data.
    from analyst.evidence import EvidenceLedger
    from analyst.runtime import AnalystReport, NumericClaim, audit_report

    ledger = EvidenceLedger(tmp_path / "ledger-audit")
    receipt = ledger.record("run_case_study", "execute_case_study(output_dir)", result)
    claims = [
        NumericClaim(
            id=c["id"],
            execution_id=receipt["execution_id"],
            path=c["path"],
            value=c["value"],
            code_sha256=receipt["code_sha256"],
        )
        for c in result["numeric_claims"]
    ]
    exported = (tmp_path / "artifacts/report.md").read_text(encoding="utf-8")
    assert exported == result["report"]
    assert audit_report(AnalystReport(report=exported, claims=claims), ledger)["verified"]
    assert not any(isinstance(c["value"], bool) for c in result["numeric_claims"])
