"""Submission controls use independent historical evidence and explicit test data."""

from __future__ import annotations

import hashlib
import json
import shutil
from pathlib import Path

import pytest
from analyst.case_study.preparation import build_snapshot
from analyst.case_study.workbook import generate_workbook
from openpyxl import load_workbook
from scripts import generate_submission as submission

ROOT = Path(__file__).resolve().parents[1]


def read_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def test_saved_canonical_submission_matches_independently_frozen_controls():
    report = read_json(ROOT / "reports/canonical_snapshot.json")
    frozen = read_json(ROOT / "evals/rebuild_independent_expected.json")
    assert report["analysis"]["metadata"]["source_sha256"] == frozen["sources"]["defects"]["sha256"]
    assert (
        report["analysis"]["metadata"]["master_sha256"]
        == frozen["sources"]["master_local"]["sha256"]
    )
    cases = frozen["independent"]["cases"]
    for identifier, result in {"canonical": report["analysis"], **report["scenarios"]}.items():
        actual, expected = result["kpis"], cases[identifier]
        assert actual["event_count"] == expected["event_count"]
        assert actual["critical_count"] == expected["critical_count"]
        assert actual["eligible_event_count"] == expected["eligible_event_count"]
        assert actual["reference_output"] == expected["reference_output"]
        assert actual["reference_intensity_per_1000"] == pytest.approx(
            expected["events_per_1000_output"], abs=1e-12
        )
        assert actual["known_repair_cost"] == expected["costs"]["known_total"]
        assert actual["cost_known_count"] == expected["costs"]["valid_count"]
        assert actual["cost_unknown_count"] == expected["costs"]["unknown_count"]
    for actual, expected in zip(
        report["analysis"]["monthly"], cases["canonical"]["monthly"], strict=True
    ):
        assert actual["month"] == expected["month"]
        assert actual["event_count"] == expected["event_count"]
        assert actual["known_repair_cost"] == expected["costs"]["known_total"]
    independent_cluster = frozen["independent"]["product_81_85_cluster"]
    cluster = report["reviewed_cohort"]
    assert cluster["product_ids"] == independent_cluster["product_ids"]
    assert cluster["kpis"]["event_count"] == independent_cluster["event_count"]
    assert cluster["kpis"]["critical_count"] == independent_cluster["severity_counts"]["Critical"]
    assert cluster["kpis"]["known_repair_cost"] == independent_cluster["costs"]["known_total"]
    assert cluster["kpis"]["reference_output"] == independent_cluster["reference_output"]
    assert {row["month"]: row["event_count"] for row in cluster["monthly"]} == independent_cluster[
        "monthly_counts"
    ]


def test_workbook_counts_distinguish_cleaning_removal_scope_and_overlapping_flags(tmp_path):
    master = b"product_id,product_name,category,production_line,monthly_output\n1,A,Category 1,Line-1,100\n"
    base = {
        "defect_id": 1,
        "product_id": 1,
        "defect_type": "Functional",
        "defect_date": "1/2/2024",
        "defect_location": "Component",
        "severity": "critical",
        "inspection_method": "Visual Inspection",
        "repair_cost": 10,
    }
    raw = [
        base,
        dict(base),
        {**base, "defect_id": 2, "severity": "Critical"},
        {**base, "defect_id": 2, "product_id": 99, "severity": "Critical"},
        {**base, "defect_id": 3, "severity": "Critical", "repair_cost": None},
        {**base, "defect_id": 4, "severity": "Critical", "defect_date": "2026-06-01"},
    ]
    data = build_snapshot(
        raw,
        {
            "source_sha256": hashlib.sha256(json.dumps(raw).encode()).hexdigest(),
            "source_url": "https://synthetic.invalid/report-test",
            "fetched_at": "2026-01-01T00:00:00Z",
        },
        master,
    )
    template = ROOT / "DATA_QUALITY_TEMPLATE.xlsx"
    before = template.read_bytes()
    generated = generate_workbook(data, tmp_path, template)
    workbook = load_workbook(generated["workbook"])
    try:
        assert [workbook["Summary"].cell(row, 2).value for row in range(4, 9)] == [6, 3, 1, 2, 2]
        assert "raw" in workbook["Summary"]["C5"].value.lower()
        assert "date-scope exclusions are separate" in workbook["Summary"]["C5"].value.lower()
        assert "source rows" in workbook["Summary"]["C6"].value
        assert "retained in-period events" in workbook["Summary"]["C7"].value
        issues = [row for row in workbook["Issues Found"].iter_rows(min_row=4) if row[2].value]
        conflict = next(row for row in issues if "contradictory" in row[2].value)
        assert conflict[5].value == "Removed" and "quarantined" in conflict[6].value
        assert all(row[5].value in {"Removed", "Corrected", "Flagged"} for row in issues)
        assert all(isinstance(row[4].value, int) for row in issues)
    finally:
        workbook.close()
    assert template.read_bytes() == before
    assert generated["workbook_delivery"]["extended_readback"][
        "action_validation_preserved_and_extended"
    ]
    with pytest.raises(ValueError, match="fixture is not a source"):
        submission.validate_real_snapshot(data)


@pytest.fixture
def delivered_copy(tmp_path, monkeypatch):
    report = read_json(ROOT / "reports/canonical_snapshot.json")
    monkeypatch.setattr(submission, "code_identity", lambda: report["code_files"])
    for name in (
        "README.md",
        "README.vi.md",
        "DATA_QUALITY.xlsx",
        "reports/canonical_snapshot.json",
        "reports/submission_manifest.json",
    ):
        target = tmp_path / name
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(ROOT / name, target)
    return tmp_path


def update_manifest_hash(directory: Path, name: str):
    manifest_path = directory / "reports/submission_manifest.json"
    manifest = read_json(manifest_path)
    manifest["files"][name] = submission.sha256((directory / name).read_bytes())
    manifest_path.write_bytes(submission.json_bytes(manifest))


@pytest.mark.parametrize(
    "fault,message", [("monthly", "monthly result differs"), ("action", "issue actions violate")]
)
def test_cross_artifact_verification_rejects_plausible_workbook_drift(
    delivered_copy, fault, message
):
    path = delivered_copy / "DATA_QUALITY.xlsx"
    workbook = load_workbook(path)
    if fault == "monthly":
        workbook["Monthly Results"]["B7"] = workbook["Monthly Results"]["B7"].value + 1
    else:
        workbook["Issues Found"]["F4"] = "Quarantined"
    workbook.save(path)
    workbook.close()
    update_manifest_hash(delivered_copy, "DATA_QUALITY.xlsx")
    with pytest.raises(ValueError, match=message):
        submission.verify(delivered_copy)


def test_cross_artifact_verification_rejects_readme_number_drift(delivered_copy):
    path = delivered_copy / "README.vi.md"
    original = path.read_text(encoding="utf-8")
    path.write_text(original.replace("| 2024-06 | 252 |", "| 2024-06 | 253 |", 1), encoding="utf-8")
    update_manifest_hash(delivered_copy, "README.vi.md")
    with pytest.raises(ValueError, match="README analysis differs"):
        submission.verify(delivered_copy)


def test_cross_artifact_verification_rejects_a_scenario_as_canonical(delivered_copy):
    path = delivered_copy / "reports/canonical_snapshot.json"
    report = read_json(path)
    report["analysis"]["scope"]["scenario"] = "exclude_cost_10"
    path.write_bytes(submission.json_bytes(report))
    update_manifest_hash(delivered_copy, "reports/canonical_snapshot.json")
    with pytest.raises(ValueError, match="canonical scope"):
        submission.verify(delivered_copy)


def test_generated_section_preserves_owner_prose_and_rejects_duplicate_markers():
    owner = f"Owner introduction\n{submission.START_MARKER}\nold\n{submission.END_MARKER}\nOwner appendix"
    generated = f"{submission.START_MARKER}\nnew\n{submission.END_MARKER}"
    assert (
        submission.replace_generated(owner, generated)
        == f"Owner introduction\n{generated}\nOwner appendix"
    )
    with pytest.raises(ValueError, match="exactly one"):
        submission.replace_generated(owner + submission.START_MARKER, generated)


@pytest.mark.parametrize("name", ["README.md", "README.vi.md"])
@pytest.mark.parametrize(
    "fault,message",
    [
        ("reversed", "ordered canonical-analysis marker pair"),
        ("missing", "exactly one canonical-analysis marker pair"),
        ("empty", "README analysis differs"),
    ],
)
def test_submission_verification_rejects_missing_canonical_region(
    delivered_copy, name, fault, message
):
    path = delivered_copy / name
    original = path.read_text(encoding="utf-8")
    if fault == "reversed":
        changed = (
            original.replace(submission.START_MARKER, "<!-- temporary start marker -->")
            .replace(submission.END_MARKER, submission.START_MARKER)
            .replace("<!-- temporary start marker -->", submission.END_MARKER)
        )
    elif fault == "missing":
        changed = original.replace(submission.START_MARKER, "").replace(submission.END_MARKER, "")
    else:
        changed = submission.replace_generated(
            original, f"{submission.START_MARKER}\n{submission.END_MARKER}"
        )
    path.write_text(changed, encoding="utf-8")
    update_manifest_hash(delivered_copy, name)
    with pytest.raises(ValueError, match=message):
        submission.verify(delivered_copy)


@pytest.mark.parametrize(
    "value,decimals,en,vi",
    [(1090, 0, "1,090", "1.090"), (1234.5, 2, "1,234.50", "1.234,50"), (None, 0, "N/A", "N/A")],
)
def test_submission_locale_formatting_preserves_value(value, decimals, en, vi):
    assert submission.number(value, "en", decimals) == en
    assert submission.number(value, "vi", decimals) == vi
