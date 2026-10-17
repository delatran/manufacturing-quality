"""Category/line mix controls, scope-aware priorities and the shareable brief."""

from __future__ import annotations

import io
import math
import re

import pytest
from analyst.case_study.analytics import Scope, analysis_report
from analyst.case_study.brief import number, render_brief
from analyst.case_study.preparation import build_snapshot
from analyst.case_study.workbook import cost_floor_issue
from analyst.dashboard import AnalysisService, register_routes
from analyst.evidence import canonical, sha256
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from tests.test_dashboard import ROOT

# Category 9 is made on both lines; each line also makes its own category.
MASTER = (
    b"product_id,product_name,category,production_line,monthly_output\n"
    b"1,Alpha,Category 1,Line-1,1000\n"
    b"2,Beta,Category 2,Line-2,1000\n"
    b"3,Gamma,Category 9,Line-1,100\n"
    b"4,Delta,Category 9,Line-2,100\n"
)


def row(identifier, product_id, month, cost, severity="Minor", location="Surface"):
    return {
        "defect_id": identifier,
        "product_id": product_id,
        "defect_type": "Cosmetic",
        "defect_date": f"2024-{month:02d}-10",
        "defect_location": location,
        "severity": severity,
        "inspection_method": "Visual Inspection",
        "repair_cost": cost,
    }


def rows():
    data = []
    identifier = 1
    plan = [
        # product, events per month (Jan..Jun), cost
        (1, (1, 1, 1, 1, 1, 1), 100.0),
        (2, (1, 1, 1, 2, 2, 2), 300.0),
        (3, (1, 1, 1, 2, 2, 2), 100.0),
        (4, (1, 1, 1, 2, 2, 2), 300.0),
    ]
    for product, counts, cost in plan:
        for month, count in enumerate(counts, start=1):
            for _ in range(count):
                data.append(
                    row(
                        identifier,
                        product,
                        month,
                        cost,
                        severity="Critical" if product == 4 else "Minor",
                        location="Component" if product == 4 else "Surface",
                    )
                )
                identifier += 1
    data.append(row(identifier, 2, 3, 10))
    return data


@pytest.fixture()
def data():
    source = rows()
    return build_snapshot(
        source,
        {
            "source_url": "https://synthetic.invalid/mix",
            "source_sha256": sha256(canonical(source)),
            "fetched_at": "2026-10-17T00:00:00Z",
        },
        MASTER,
    )


def test_shared_category_is_compared_with_the_rest_of_each_line(data):
    report = analysis_report(data, Scope())
    summary = {item["category"]: item for item in report["category_summary"]}
    assert summary["Category 9"]["production_lines"] == ["Line-1", "Line-2"]
    assert summary["Category 9"]["eligible_event_count"] == 18
    assert summary["Category 9"]["reference_output"] == 1200
    assert math.isclose(summary["Category 9"]["reference_intensity_per_1000"], 15.0)
    assert math.isclose(summary["Category 9"]["eligible_events_per_cohort_product"], 9.0)
    assert math.isclose(summary["Category 9"]["average_monthly_output_per_product"], 100.0)
    assert report["category_summary"][0]["category"] == "Category 9"
    control = next(
        item for item in report["shared_category_controls"] if item["category"] == "Category 9"
    )
    assert control["lines_compared"] == control["lines_above_rest_of_line"] == 2
    by_line = {item["production_line"]: item for item in control["lines"]}
    assert math.isclose(by_line["Line-1"]["reference_intensity_per_1000"], 15.0)
    assert by_line["Line-1"]["rest_of_line_categories"] == ["Category 1"]
    assert math.isclose(by_line["Line-1"]["rest_of_line_reference_intensity_per_1000"], 1.0)
    assert math.isclose(by_line["Line-1"]["intensity_ratio_to_rest_of_line"], 15.0)
    assert math.isclose(by_line["Line-2"]["cost_ratio_to_same_category_other_lines"], 3.0)


def test_line_cost_comparison_and_half_periods_reconcile(data):
    report = analysis_report(data, Scope())
    costs = report["line_cost_comparison"]
    assert costs[0]["production_line"] == "Line-2"
    total = sum(item["known_repair_cost"] for item in costs)
    assert math.isclose(total, report["kpis"]["known_repair_cost"])
    assert math.isclose(sum(item["known_cost_share_pct"] for item in costs), 100.0)
    halves = report["period_halves"]
    assert halves["enabled"]
    assert halves["first"]["months"] == ["2024-01", "2024-02", "2024-03"]
    assert (
        halves["first"]["event_count"] + halves["second"]["event_count"]
        == (report["kpis"]["event_count"])
    )
    assert analysis_report(data, Scope(end_month="2024-04"))["period_halves"]["enabled"]
    assert not analysis_report(data, Scope(end_month="2024-05"))["period_halves"]["enabled"]


def test_priorities_follow_severity_frequency_and_cost_with_scope_fallbacks(data):
    report = analysis_report(data, Scope())
    kinds = [item["kind"] for item in report["recommendations"]]
    assert kinds == ["critical_location", "category_intensity", "line_repair_cost"]
    category = report["recommendations"][1]
    assert category["filters"] == {"category": "Category 9"}
    assert category["metrics"]["lines_above_rest_of_line"] == 2
    assert math.isclose(
        category["metrics"]["rest_of_scope_reference_intensity_per_1000"], 16 / 12000 * 1000
    )
    cost = report["recommendations"][2]
    assert cost["metrics"]["production_line"] == "Line-2"
    assert cost["metrics"]["ratio_to_other_lines"] >= 1.2
    single_category = analysis_report(data, Scope(category="Category 1"))
    assert [item["kind"] for item in single_category["recommendations"]] == [
        "line_intensity",
        "defect_type_cost",
    ]
    single_line = analysis_report(data, Scope(production_line="Line-1"))
    assert [item["kind"] for item in single_line["recommendations"]] == [
        "category_intensity",
        "defect_type_cost",
    ]
    for item in single_line["recommendations"]:
        assert item["metrics"] and item["title"] and item["evidence"]
    empty = analysis_report(data, Scope(severity="Critical", production_line="Line-1"))
    assert empty["recommendations"] == []


def test_trend_interpretation_states_the_direction_of_the_signal(data):
    text = analysis_report(data, Scope())["trend"]["interpretation"]
    assert "adverse signal" in text and text.endswith("defect-rate change.")


@pytest.mark.parametrize("language", ["en", "vi"])
def test_brief_reformats_backend_numbers_without_inventing_values(data, language):
    payload = {
        "context": {"context_id": "c" * 64},
        "analysis": analysis_report(data, Scope()),
    }
    text = render_brief(payload, language)
    kpis = payload["analysis"]["kpis"]
    assert number(kpis["event_count"], language) in text
    assert number(kpis["known_repair_cost"], language, 2) in text
    assert not re.search(r"\bNone\b|\bnan\b|[{}]", text)
    headings = [line for line in text.splitlines() if line.startswith("### ")]
    assert len(headings) == len(payload["analysis"]["recommendations"]) == 3
    if language == "vi":
        assert text.startswith("# Bản tin chất lượng")
        assert "Investigate" not in text and "Evidence" not in text
    else:
        assert text.startswith("# Quality brief: whole plant")


def test_brief_endpoint_serves_the_context_in_both_languages(tmp_path, data):
    subject = AnalysisService(ROOT, tmp_path / "state", snapshot_loader=lambda: data)
    app = FastAPI()
    register_routes(app, subject)
    client = TestClient(app)
    payload = client.get("/api/dashboard").json()
    link = payload["links"]["brief"]
    english = client.get(link)
    assert english.status_code == 200
    assert english.headers["content-type"].startswith("text/markdown")
    assert "attachment" in english.headers["content-disposition"]
    assert "Quality brief" in english.text
    vietnamese = client.get(link, params={"locale": "vi"})
    assert vietnamese.headers["content-language"] == "vi"
    assert "Bản tin chất lượng" in vietnamese.text
    assert client.get(link, params={"locale": "fr"}).status_code == 422
    assert client.get(link, params={"other": "1"}).status_code in {400, 422}


def test_cost_floor_issue_is_documented_without_changing_canonical_costs(data):
    issue = cost_floor_issue(data, data.quality)
    assert issue[0] == "Defects / repair_cost"
    assert issue[3] == data.quality["raw_profile"]["cost_ten_rows"] == 1
    assert issue[4] == "Flagged"
    assert "exclude_cost_10" in issue[5]
    report = analysis_report(data, Scope())
    assert any(
        item["production_line"] == "Line-2" and item["known_repair_cost"] > 0
        for item in report["line_cost_comparison"]
    )


def test_workbook_lists_specific_detection_rules_and_the_cost_floor(tmp_path, data):
    from analyst.case_study.workbook import generate_workbook

    generate_workbook(data, tmp_path, ROOT / "DATA_QUALITY_TEMPLATE.xlsx")
    workbook = load_workbook(io.BytesIO((tmp_path / "DATA_QUALITY.xlsx").read_bytes()))
    sheet = workbook["Issues Found"]
    rows = [
        [cell.value for cell in line]
        for line in sheet.iter_rows(min_row=4)
        if line[0].value is not None
    ]
    detections = [line[3] for line in rows]
    assert not any(str(text).startswith("Explicit domain") for text in detections)
    floor = [line for line in rows if str(line[2]).startswith("A cluster of repair costs")]
    assert len(floor) == 1 and floor[0][5] == "Flagged"
