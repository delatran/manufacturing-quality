"""Regenerate the completed workbook and empirical summary from the real HTTP pipeline.

Run from the project root:
    python -m analyst.case_study --output artifacts/case-study

Only the master CSV and Excel template are local inputs. No raw defect JSON
is saved. The supplied template is never overwritten.
"""

import hashlib
import io
import json
import math
import statistics
import textwrap
from collections.abc import Iterable, Sequence
from copy import copy
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Any, cast

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection
from openpyxl.styles.fills import Fill
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from .analytics import DEFINITIONS, Scope, analysis_report
from .config import MONTHS
from .models import AuditEntry, Snapshot
from .scenarios import SCENARIOS


def cost_floor_issue(snapshot: Snapshot, quality: dict[str, Any]) -> list[Any]:
    """Describe the repeated minimum repair cost without changing canonical values."""
    floor = SCENARIOS["exclude_cost_10"]["parameters"]["repair_cost_equal"]
    raw_rows = quality.get("raw_profile", {}).get("cost_ten_rows", 0)
    raw_values = [json.loads(text).get("repair_cost") for text in snapshot.raw_records]
    integer_rows = sum(type(value) is int and value == floor for value in raw_values)
    in_period = [
        event for event in snapshot.events if event.defect_date.strftime("%Y-%m") in MONTHS
    ]
    floor_costs = [
        event.repair_cost
        for event in in_period
        if event.repair_cost is not None and event.repair_cost == floor
    ]
    valid = [event.repair_cost for event in in_period if event.repair_cost is not None]
    stored = (
        f" ({integer_rows} stored as integers while other observed costs are decimals)"
        if raw_values
        else ""
    )
    median = (
        f"; the in-period median valid cost is {statistics.median(valid):,.2f}" if valid else ""
    )
    return [
        "Defects / repair_cost",
        f"A cluster of repair costs equals exactly {floor}, the lowest observed value. It may be a minimum charge or a placeholder rather than a measured repair cost.",
        f"Value-frequency profile of repair_cost: {raw_rows} raw rows equal {floor}{stored}{median}.",
        raw_rows,
        "Flagged",
        f"Count basis: received raw defect rows. Kept as valid observed cost because no source rule marks it invalid; {len(floor_costs)} retained in-period events carry it. "
        f"Scenario exclude_cost_10 withholds these costs, lowering known cost by {math.fsum(floor_costs):,.2f} without changing event counts. "
        "Not an audit flag, so it is outside the Summary flagged count. Confirm the cost-entry rule with Finance.",
    ]


def safe_cell(value: Any) -> Any:
    """Keep source-derived text as text in Excel, including formula prefixes."""
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def write_row(sheet: Worksheet, row: int, values: Sequence[Any], *, template_row: int = 4) -> None:
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row, column)
        if row > sheet.max_row or not cell.has_style:
            reference = sheet.cell(template_row, column)
            cell.font = cast(Font, copy(reference.font))
            cell.fill = cast(Fill, copy(reference.fill))
            cell.border = cast(Border, copy(reference.border))
            cell.protection = cast(Protection, copy(reference.protection))
            cell.number_format = reference.number_format
        cell.value = safe_cell(value)
        cell.alignment = Alignment(vertical="top", wrap_text=True)


def add_table(
    workbook: Workbook, name: str, headers: list[str], rows: Sequence[Sequence[Any]]
) -> Worksheet:
    sheet = workbook.create_sheet(name)
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="1C54D4")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        sheet.append([safe_cell(value) for value in row])
    sheet.freeze_panes = "A2"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        width = min(65, max(14, max(len(str(cell.value or "")) for cell in column[:30]) + 2))
        sheet.column_dimensions[letter].width = width
        for cell in column[1:]:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    return sheet


def issue_action(entries: Iterable[AuditEntry], code: str) -> str:
    """Map policy decisions onto the template's permitted action vocabulary."""
    if code in {
        "exact_duplicate",
        "normalized_duplicate",
        "complementary_duplicate",
        "conflicting_id",
        "invalid_identifier",
        "invalid_date",
        "master_duplicate",
    }:
        return "Removed"
    if any(entry.action == "corrected" for entry in entries):
        return "Corrected"
    return "Flagged"


def generate_workbook(
    snapshot: Snapshot,
    output_dir: Path,
    template_path: Path,
    *,
    analysis: dict[str, Any] | None = None,
    report_snapshot: dict[str, Any] | None = None,
) -> dict[str, Any]:
    output = output_dir / "DATA_QUALITY.xlsx"
    if output.resolve() == template_path.resolve() or (
        output.exists() and output.samefile(template_path)
    ):
        raise ValueError("Completed workbook output must be separate from the supplied template")
    template_bytes = template_path.read_bytes()
    template_sha256 = hashlib.sha256(template_bytes).hexdigest()
    analysis = analysis if analysis is not None else analysis_report(snapshot, Scope())
    quality = snapshot.quality
    if (
        analysis["scope"] != asdict(Scope())
        or analysis["kpis"]["event_count"] != quality["in_scope_event_count"]
    ):
        raise ValueError("The quality workbook requires the complete canonical analysis scope")
    if analysis["metadata"]["snapshot_id"] != snapshot.provenance["snapshot_id"]:
        raise ValueError("Workbook analysis and source snapshot identities differ")
    workbook = load_workbook(io.BytesIO(template_bytes))
    expected_sheets = ["Summary", "Issues Found", "Assumptions", "Impact on Results"]
    if workbook.sheetnames != expected_sheets:
        raise ValueError(
            "Unexpected template sheet layout; refusing to write a mismatched workbook"
        )
    summary = workbook["Summary"]
    summary.cell(3, 3, "Counting basis").font = copy(summary["B3"].font)
    summary.cell(3, 3).fill = copy(summary["B3"].fill)
    summary.cell(3, 3).border = copy(summary["B3"].border)
    summary_rows = [
        (4, quality["raw_event_count"], "HTTP payload rows, before any cleaning."),
        (
            5,
            quality["reconciliation"]["duplicate_copies_removed"] + quality["quarantined_count"],
            f"Raw source rows removed during cleaning: {quality['reconciliation']['duplicate_copies_removed']} surplus copies and {quality['quarantined_count']} quarantined variants. Date-scope exclusions are separate below.",
        ),
        (
            6,
            quality["corrected_source_rows"],
            "Distinct source rows with case/whitespace or unambiguous date-format correction. Overlaps other categories.",
        ),
        (
            7,
            quality["flagged_in_scope_event_count"],
            "Distinct retained in-period events with at least one audit flag, including convention-dependent dates. Flags overlap and are not removed.",
        ),
        (
            8,
            quality["in_scope_event_count"],
            "January-June 2024 unique nonconflicting events, including unmatched products and unknown costs.",
        ),
    ]
    for row, count, note in summary_rows:
        summary.cell(row, 2, count)
        summary.cell(row, 3, note).alignment = Alignment(wrap_text=True, vertical="top")
        summary.row_dimensions[row].height = 46
    summary.column_dimensions["C"].width = 100
    extras = [
        (
            "Master records received",
            quality["master_raw_count"],
            "Raw master CSV rows; counted separately from received HTTP defect rows.",
        ),
        (
            "Master products retained",
            quality["master_clean_count"],
            "Validated unique positive product IDs, complete dimensions and positive integer monthly_output.",
        ),
        (
            "Exact duplicate surplus",
            quality["exact_duplicate_count"],
            "Byte-equivalent object payload copies after JSON key-order canonicalization.",
        ),
        (
            "Normalized duplicate surplus",
            quality["normalized_duplicate_count"],
            "Surplus source rows equivalent after documented representation normalization.",
        ),
        (
            "Complementary cost duplicate surplus",
            quality["complementary_duplicate_count"],
            "Surplus source rows with one missing and one observed cost, with all other event attributes matching.",
        ),
        (
            "Conflicting variants quarantined",
            quality["quarantined_count"],
            "Raw source rows excluded from clean events, without choosing an unverified winning variant.",
        ),
        (
            "Clean events before date scope",
            quality["clean_event_count"],
            "Unique retained events after duplicate and conflict resolution, across all actual source years.",
        ),
        (
            "Events outside the study period",
            quality["out_of_scope_count"],
            "Retained clean events excluded only by the analytical date scope, without changing their year.",
        ),
        (
            "Linked events used in reference-output intensity",
            quality["linked_event_count"],
            "Retained in-period events matched to one master product.",
        ),
        (
            "Unmatched in-period events",
            quality["unmatched_event_count"],
            "Retained in plant counts and costs, excluded from linked intensity without a matching output denominator.",
        ),
        (
            "Events with known repair cost",
            quality["known_cost_count"],
            "Valid observed nonnegative cost; no imputation.",
        ),
        (
            "Events with unknown repair cost",
            quality["unknown_cost_count"],
            "Retained in-period events; missing, invalid and negative costs remain unknown without imputation.",
        ),
        (
            "Calendar dates without logged events",
            quality["temporal_coverage"]["days_without_events"],
            "Calendar dates in the study period; operating/reporting coverage is unverified.",
        ),
    ]
    for row, extra_values in enumerate(extras, start=10):
        write_row(summary, row, extra_values)
        summary.row_dimensions[row].height = 42
    for summary_cells in summary.iter_rows(min_row=4, max_col=3):
        for cell in summary_cells:
            cell.alignment = Alignment(
                horizontal=cell.alignment.horizontal, vertical="top", wrap_text=True
            )
        if isinstance(summary_cells[1].value, (int, float)) and not isinstance(
            summary_cells[1].value, bool
        ):
            summary_cells[1].number_format = "#,##0"

    issues_sheet = workbook["Issues Found"]
    analytical_rationale = {
        "ambiguous_date_assumption": "A day/month reversal can move events between months and change the trend. Declaring the convention makes that uncertainty visible pending source-owner confirmation.",
        "complementary_duplicate": "Retaining the one observed cost preserves supported financial evidence while preventing duplicate event counts and avoiding cost imputation.",
        "conflicting_id": "Selecting one contradictory variant would arbitrarily assign the event and its cost to a product or production line, biasing comparisons.",
        "day_first_date": "An unambiguous valid date can be recovered without guessing. Excluding it would unnecessarily undercount its month, while the audit preserves the format exception.",
        "exact_duplicate": "Repeated copies would inflate event counts and repair-cost totals without representing additional defects.",
        "label_normalized": "Canonical labels keep the same defect or severity category together, preventing fragmented groups and misleading category shares.",
        "missing_cost": "Dropping the event would undercount defects, while zero or imputed cost would misstate the financial burden. Explicit coverage shows the unresolved cost gap.",
        "negative_cost": "A negative cost-incurred value would reduce the repair subtotal without verified credit or reversal evidence. The event remains countable while its cost awaits correction.",
        "normalized_duplicate": "Formatting differences do not establish a second defect event. Collapsing equivalent records avoids inflated counts after normalization.",
        "outside_study_period": "Changing a valid source year to fit the requested period would invent evidence. Scope exclusion preserves the original date; alternate years belong to an explicit scenario.",
        "unmatched_product": "An inner join would hide these events from plant totals. Keeping them unassigned preserves burden while avoiding an unsupported line/category assignment or output denominator.",
        "calendar_reporting_coverage": "Treating absent records as verified zero-defect production would overstate completeness and distort daily screening. Reporting and operating calendars are needed to interpret the gaps.",
        "invalid_cost": "An unsupported cost value cannot support a financial sum. Retaining the event with unknown cost prevents both defect undercounting and invented monetary values.",
        "invalid_date": "An impossible or unsupported date cannot be assigned to a reporting period reliably; quarantine prevents fabricated temporal patterns.",
        "invalid_identifier": "An invalid key cannot establish a unique event or a trustworthy product join; quarantine avoids invented identity or assignment.",
        "unexpected_label": "Mapping an unrecognized value to a named category would invent its meaning. An explicit Unknown group preserves counts and exposes the source correction needed.",
        "missing_label": "An absent category cannot support a specific assignment. Keeping an Unknown group avoids dropping events or overstating known-category coverage.",
        "invalid_output": "Using partial or invalid production exposure would bias normalized comparisons; associated cohort intensity remains unavailable until the denominator is verified.",
        "master_duplicate": "Duplicate master rows would multiply joined events and reference output; a unique product dimension prevents join inflation.",
        "master_text_normalized": "Whitespace-only label differences would split otherwise identical product cohorts and complicate consistent filtering.",
    }
    for row, issue in enumerate(quality["issues"], start=4):
        entries = [entry for entry in snapshot.audit if entry.code == issue["code"]]
        fields = sorted({entry.field for entry in entries})
        detection = {
            "exact_duplicate": "Canonical JSON equality over every field, preserving first source-row lineage.",
            "normalized_duplicate": "Group by event ID after controlled vocabulary normalization, compare all event attributes.",
            "complementary_duplicate": "Same event ID and all non-cost attributes; only one distinct observed cost and a null counterpart.",
            "conflicting_id": "Same event ID has more than one non-equivalent attribute vector after normalization.",
            "ambiguous_date_assumption": "Slash-date components are both at most twelve and unequal; date order requires a declared assumption. Equal components do not change under reversal.",
            "calendar_reporting_coverage": "Compare retained event dates with every calendar date in the study period; inspect month-end maximum dates.",
            "day_first_date": "Slash dates whose first component exceeds twelve cannot be month-first; the second component is a valid month, so the date is read as D/M/YYYY.",
            "label_normalized": "Value counts per label field: variants with leading/trailing spaces or different letter case match a known domain value only after trimming and case-insensitive comparison.",
            "missing_cost": "repair_cost is null or blank in the payload.",
            "negative_cost": "repair_cost parses as a number below zero, which a cost-incurred field cannot hold.",
            "outside_study_period": "Parsed event month falls outside January-June 2024 after duplicate resolution; the year is kept as supplied.",
            "unmatched_product": "Left join of event product_id to the validated master product_id finds no master row.",
            "invalid_cost": "repair_cost is present but not a finite number in a supported format.",
            "invalid_date": "defect_date is missing or matches no supported ISO or slash format with a valid calendar date.",
            "invalid_identifier": "defect_id or product_id is missing, non-integer or not positive.",
            "unexpected_label": "Label value does not match the controlled vocabulary even after trimming and case-insensitive comparison.",
            "missing_label": "Label field is null or blank.",
            "invalid_output": "monthly_output is missing, non-integer or not positive.",
            "master_duplicate": "product_id appears more than once in the master CSV.",
            "master_text_normalized": "Master text fields differ only by surrounding whitespace.",
        }.get(
            issue["code"],
            "Explicit domain, range, format or foreign-key check in analyst/case_study/preparation.py; row evidence is in Audit Detail.",
        )
        affected_unit = issue.get("affected_unit", "source rows audited after exact-copy removal")
        if issue["code"] in {"unmatched_product", "outside_study_period"}:
            affected_unit = "retained clean events before the study-period filter"
        if issue["code"] == "exact_duplicate":
            affected_unit = "surplus raw source rows"
        detail = f"Count basis: {affected_unit}. {issue['handling']} {analytical_rationale.get(issue['code'], '')}"
        if issue["code"] == "conflicting_id":
            detail += " Template action Removed means quarantined from analysis; original source evidence remains unchanged."
        if issue["code"] == "outside_study_period":
            detail += " Flagged means retained in the clean snapshot and excluded only by analytical date scope."
        write_row(
            issues_sheet,
            row,
            [
                row - 3,
                f"{issue['source']} / {', '.join(fields) or 'calendar coverage'}",
                issue["description"],
                detection,
                issue["affected_count"],
                issue_action(entries, issue["code"]),
                detail,
            ],
        )
        issues_sheet.row_dimensions[row].height = 96
    additional_issues = [
        [
            "Master / monthly_output",
            "Only static reference monthly output is available, without actual monthly or daily exposure.",
            "Master schema has one monthly_output value per product and no month/actual-output field.",
            quality["master_clean_count"],
            "Flagged",
            "Count basis: retained master products. Use all eligible master products × selected full months; label reference-output intensity. Logged-event counts are not defective-unit probabilities.",
        ],
        [
            "Master / category and production_line",
            "Line and category allocations are interdependent, so total line rankings mix effects.",
            "Cross-tabulate unique master products by category and line.",
            quality["master_clean_count"],
            "Flagged",
            "Count basis: retained master products. Compare category/product mix and actual exposure before attributing a difference to line performance.",
        ],
        [
            "Defects / repair_cost unit",
            "Local currency name, price basis and cost scope are unspecified.",
            "Brief specifies local currency only; payload has no currency or cost-component field.",
            quality["raw_event_count"],
            "Flagged",
            "Count basis: received raw defect rows. Keep unspecified local currency, observed cost only and explicit coverage; no currency conversion or unobserved cost estimate.",
        ],
        [
            "Defects / granularity",
            "No unit-level product identifier or inspected-unit count is supplied.",
            "Schema contains event ID and product SKU, without serial/unit/batch ID or inspection denominator.",
            quality["raw_event_count"],
            "Flagged",
            "Count basis: received raw defect rows. Describe logged defect events. Repeated events on one unit and detection changes cannot be separated; defective-unit probability, yield and causal improvement are unavailable.",
        ],
        [
            "Source / freshness",
            "Historical event dates and the HTTP retrieval timestamp describe different periods.",
            "Compare source event dates, study bounds and HTTP retrieval timestamp.",
            quality["raw_event_count"],
            "Flagged",
            "Count basis: received raw defect rows. Display event scope and retrieval time separately. A current fetch timestamp does not establish current production coverage.",
        ],
        cost_floor_issue(snapshot, quality),
    ]
    if not additional_issues[-1][3]:
        additional_issues.pop()
    for row, issue_values in enumerate(additional_issues, start=4 + len(quality["issues"])):
        write_row(issues_sheet, row, [row - 3, *issue_values])
        issues_sheet.row_dimensions[row].height = 96
    final_issue_row = 3 + len(quality["issues"]) + len(additional_issues)
    for row in range(final_issue_row + 1, issues_sheet.max_row + 1):
        for cell in issues_sheet[row]:
            cell.value = None
    for validation in issues_sheet.data_validations.dataValidation:
        if validation.type == "list" and validation.formula1 == '"Removed,Corrected,Flagged"':
            validation.sqref = f"F4:F{max(23, final_issue_row)}"

    assumptions = [
        "Primary analytical scope is January 1-June 30, 2024, with inclusive whole-month filters. Actual source years remain unchanged; valid out-of-period events are retained in the clean snapshot without assuming a typo.",
        "A defect_id identifies one logged event. Exact copies collapse; normalized equivalents collapse; contradictory variants are all quarantined.",
        "Complementary null/observed cost rows represent the same event only when all other attributes agree; retain the observed cost without imputation. Contradictory product identity is not resolved by a successful master join.",
        f"M/D/YYYY is the declared default slash-date convention, supported by unambiguous source dates. {quality['raw_profile']['ambiguous_date_rows']} order-sensitive raw slash dates require source-owner confirmation. Equal day/month components do not change under reversal. ISO and unambiguous D/M exceptions are parsed explicitly.",
        "Missing, invalid, nonfinite, ambiguous-format and negative costs are unknown. Valid zero costs are preserved. Observed costs equal to 10 remain valid; the exclude_cost_10 scenario tests withholding them. Sum/mean use observed valid costs only, with coverage.",
        "Master product-to-line/category mapping is static; product reassignment over time cannot be verified from these sources.",
        "monthly_output is a repeated reference output per product/month, not an observed production history. Include zero-event products; only product/line/category filters change the denominator cohort.",
        "Unmatched product IDs remain in plant totals under Unknown line/category, but cannot contribute to linked reference-output intensity or master-product Pareto totals.",
        "Currency remains 'local currency (unspecified)'; no exchange rate, currency name or cost inflation adjustment is inferred.",
        "Calendar dates without source events mean zero logged events, not verified zero defects or non-operating days. Source timestamps alone do not establish reporting completeness or an operating calendar.",
        "The January daily baseline is frozen at 31 calendar days and screened against February-June with identical non-date filters. Mean + 3√mean candidates are exploratory; equal exposure, independence, Poisson behavior and stable inspection are unverified.",
        "Chi-square is exploratory and conditional on valid expected counts. Cramer's V is an association size, not a causal claim; repeated units and inspection mix are unavailable.",
        "Issue counts overlap. Summary removed counts cover cleaning removals in raw rows; valid dates outside the study period are shown separately. Corrected counts use distinct raw rows; flagged-but-kept counts use retained in-period events. These subsets are non-additive.",
        "Products are ranked independently by event count, known repair cost and reference-output intensity. Count/cost concentration is reported without assuming an 80/20 rule; intensity is not additive and has no cumulative share.",
    ]
    for row, text in enumerate(assumptions, start=4):
        write_row(workbook["Assumptions"], row, [row - 3, text])
        workbook["Assumptions"].row_dimensions[row].height = 62

    impacts = [
        [
            "January-June event rows",
            quality["in_scope_event_count"],
            snapshot.raw_summary["in_period_raw_count"],
            "Same Jan-Jun scope and declared date convention. Raw includes duplicate/conflicting payload rows; all valid out-of-period rows are excluded from both columns.",
        ],
        [
            "January-June observed numeric cost",
            quality["known_repair_cost"],
            snapshot.raw_summary["raw_numeric_cost_including_negatives"],
            "Raw numeric sum skips nulls but includes negatives and duplicated/conflicting rows. Clean subtotal uses only valid observed costs; neither is an estimate for missing costs.",
        ],
        [
            "Clean events before analytical date scope",
            quality["clean_event_count"],
            quality["raw_event_count"],
            "Both columns include all source years. With cleaning uses unique retained events; without cleaning uses received raw rows. Difference is cleaning removals only; valid out-of-period dates are not a correction.",
        ],
        [
            "Master products retained",
            quality["master_clean_count"],
            quality["master_raw_count"],
            "Same complete product-master scope. Validate identifiers, dimensions and reference output before joining; never sum output after an event join.",
        ],
    ]
    for month in analysis["monthly"]:
        impacts.append(
            [
                f"Events {month['month']}",
                month["event_count"],
                snapshot.raw_summary["monthly_counts"].get(month["month"], 0),
                "Same month boundaries and date convention; all event severities, methods and product linkage statuses included.",
            ]
        )
    for row, impact_values in enumerate(impacts, start=4):
        write_row(workbook["Impact on Results"], row, impact_values)
        workbook["Impact on Results"].row_dimensions[row].height = 66
        for column in (2, 3):
            cell = workbook["Impact on Results"].cell(row, column)
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                cell.number_format = (
                    "#,##0.00"
                    if impact_values[0] == "January-June observed numeric cost"
                    else "#,##0"
                )

    audit_dicts = [asdict(entry) for entry in snapshot.audit]
    audit_headers = (
        list(audit_dicts[0])
        if audit_dicts
        else [
            "source",
            "source_row",
            "record_id",
            "code",
            "severity",
            "field",
            "raw_value",
            "action",
            "detail",
        ]
    )
    add_table(
        workbook,
        "Audit Detail",
        audit_headers,
        [
            [
                json.dumps(entry[key], ensure_ascii=False)
                if isinstance(entry[key], (dict, list, tuple))
                else entry[key]
                for key in audit_headers
            ]
            for entry in audit_dicts
        ],
    )
    add_table(
        workbook,
        "Reconciliation",
        ["Metric", "Value", "Definition"],
        [
            [
                key,
                value,
                "Raw = removed duplicate copies + quarantined + clean; clean = in-period + out-of-period.",
            ]
            for key, value in quality["reconciliation"].items()
        ],
    )
    add_table(
        workbook,
        "Provenance",
        ["Field", "Value"],
        [[key, str(value)] for key, value in snapshot.provenance.items()]
        + [
            ["Template SHA-256", template_sha256],
            ["Canonical report ID", (report_snapshot or {}).get("report_id", "Not supplied")],
            [
                "Canonical report schema",
                (report_snapshot or {}).get("report_schema_version", "Not supplied"),
            ],
        ]
        + [
            [
                f"{name} / {field}",
                json.dumps(value, ensure_ascii=False)
                if isinstance(value, (dict, list))
                else str(value),
            ]
            for name, receipt in (report_snapshot or {}).get("source_receipts", {}).items()
            for field, value in receipt.items()
        ]
        + [
            [
                "HTTP-only rule",
                "Defect records fetched into memory via httpx. No local defect JSON input or cache.",
            ],
            ["Reproduction command", "python -m analyst.case_study --output artifacts/case-study"],
            [
                "Workbook scope",
                "Full January-June 2024 case and complete source audit; dashboard filters do not change this workbook. Source identity and retrieval date appear above.",
            ],
            ["Audit stage", quality["audit_stage_note"]],
        ],
    )
    add_table(
        workbook,
        "Metric Definitions",
        ["Metric", "Definition"],
        [[key, value] for key, value in DEFINITIONS.items()],
    )
    monthly_sheet = add_table(
        workbook,
        "Monthly Results",
        [
            "Month",
            "Events",
            "Critical",
            "Critical share %",
            "Known cost",
            "Known-cost events",
            "Unknown-cost events",
        ],
        [
            [
                row["month"],
                row["event_count"],
                row["critical_count"],
                row["critical_share_pct"],
                row["known_repair_cost"],
                row["cost_known_count"],
                row["cost_unknown_count"],
            ]
            for row in analysis["monthly"]
        ],
    )
    for monthly_cells in monthly_sheet.iter_rows(min_row=2):
        for cell in monthly_cells[1:]:
            cell.number_format = "#,##0.00" if cell.column in (4, 5) else "#,##0"
    calendar_sheet = add_table(
        workbook,
        "Calendar Gaps",
        ["Date", "Interpretation"],
        [
            [
                datetime.fromisoformat(value),
                "No logged events; operating/reporting completeness is unknown.",
            ]
            for value in quality["temporal_coverage"]["dates_without_events"]
        ],
    )
    for calendar_cells in calendar_sheet.iter_rows(min_row=2):
        calendar_cells[0].number_format = "yyyy-mm-dd"
    for sheet in workbook:
        if sheet.title in expected_sheets:
            sheet.freeze_panes = "B4"
            sheet.sheet_view.showGridLines = False
            # Preserve the supplied column widths and header styles. Expand row
            # height for wrapped prose so long decisions remain readable in Excel.
        for sheet_cells in sheet.iter_rows(min_row=4 if sheet.title in expected_sheets else 2):
            required_height = 0
            for cell in sheet_cells:
                if cell.value is None:
                    continue
                width = sheet.column_dimensions[cell.column_letter].width or 13
                wrap_width = max(3, int(width * 1.1))
                display_value = (
                    cell.value.strftime("%Y-%m-%d")
                    if isinstance(cell.value, datetime) and cell.number_format == "yyyy-mm-dd"
                    else str(cell.value)
                )
                line_count = sum(
                    max(1, len(textwrap.wrap(line, width=wrap_width)))
                    for line in display_value.split("\n")
                )
                font_size = cell.font.sz or 11
                required_height = max(required_height, math.ceil(line_count * font_size * 1.3 + 8))
            row_index = sheet_cells[0].row
            if required_height and row_index is not None:
                sheet.row_dimensions[row_index].height = min(409, max(required_height, 20))
    expected_cells = {
        (sheet.title, cell.coordinate): cell.value
        for sheet in workbook
        for row in sheet
        for cell in row
        if cell.value not in (None, "")
    }
    workbook.save(output)
    workbook.close()
    # Bind the checksum and finite read-back checks to the same saved bytes.
    output_bytes = output.read_bytes()
    output_sha256 = hashlib.sha256(output_bytes).hexdigest()
    verified = load_workbook(io.BytesIO(output_bytes), data_only=False)
    try:
        observed_sheets = verified.sheetnames
        observed_analysis_events = (
            verified["Summary"]["B8"].value if "Summary" in observed_sheets else None
        )

        def equal_saved(actual: Any, expected: Any) -> bool:
            if isinstance(expected, float) and isinstance(actual, (int, float)):
                return math.isclose(actual, expected, rel_tol=1e-14, abs_tol=1e-12)
            return actual == expected

        mismatches = [
            f"{sheet}!{address}"
            for (sheet, address), value in expected_cells.items()
            if not equal_saved(verified[sheet][address].value, value)
        ]
        action_rules = verified["Issues Found"].data_validations.dataValidation
        valid_action_rule = any(
            rule.type == "list"
            and rule.formula1 == '"Removed,Corrected,Flagged"'
            and all(f"F{row}" in rule for row in range(4, final_issue_row + 1))
            for rule in action_rules
        )
        saved_actions = [
            verified["Issues Found"].cell(row, 6).value for row in range(4, final_issue_row + 1)
        ]
    finally:
        verified.close()
    if (
        observed_sheets[:4] != expected_sheets
        or observed_analysis_events != quality["in_scope_event_count"]
    ):
        raise AssertionError("Completed workbook read-back failed")
    if (
        mismatches
        or not valid_action_rule
        or any(action not in {"Removed", "Corrected", "Flagged"} for action in saved_actions)
    ):
        raise AssertionError(
            f"Completed workbook extended read-back failed: {mismatches[:5]}; valid action rule={valid_action_rule}"
        )
    template_after_sha256 = hashlib.sha256(template_path.read_bytes()).hexdigest()
    if template_after_sha256 != template_sha256:
        raise AssertionError("Supplied workbook template changed during generation")
    delivery = {
        "mode": "completed_copy_from_supplied_template",
        "template": {
            "name": template_path.name,
            "sha256_before": template_sha256,
            "sha256_after": template_after_sha256,
            "preserved": True,
        },
        "output": {"name": output.name, "sha256": output_sha256},
        "readback": {
            "status": "verified",
            "scope": "first_four_sheet_names_and_summary_B8",
            "expected_first_sheet_names": expected_sheets,
            "observed_sheet_names": observed_sheets,
            "checked_cell": "Summary!B8",
            "expected_analysis_events": quality["in_scope_event_count"],
            "observed_analysis_events": observed_analysis_events,
            "limitation": "Only the first four sheet names in order and Summary!B8 were checked; other cells, formatting and visual layout were not certified.",
        },
        "extended_readback": {
            "status": "verified",
            "populated_cells_checked": len(expected_cells),
            "issue_actions_checked": len(saved_actions),
            "action_validation_preserved_and_extended": True,
            "visual_layout": "Not run by the portable generator; rendered verification is separate.",
        },
    }
    summary_path = output_dir / "empirical_summary.json"
    summary_path.parent.mkdir(parents=True, exist_ok=True)
    summary_path.write_text(
        json.dumps(
            {**analysis, "raw_comparison": snapshot.raw_summary},
            indent=2,
            ensure_ascii=False,
            allow_nan=False,
        )
        + "\n",
        encoding="utf-8",
    )
    return {
        "workbook": str(output),
        "summary": str(summary_path),
        "snapshot_id": snapshot.provenance["snapshot_id"],
        "events": quality["in_scope_event_count"],
        "known_cost": quality["known_repair_cost"],
        "workbook_delivery": delivery,
    }
