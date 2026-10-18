"""Generate bilingual submission analysis and the workbook from one HTTP snapshot.

Run from the repository root. --check verifies the saved snapshot, local inputs,
current calculation code and deliverable read-back without fetching a new source.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from dataclasses import asdict
from datetime import UTC, datetime
from pathlib import Path
from tempfile import TemporaryDirectory
from typing import Any

import httpx
from analyst.case_study.analytics import (
    Scope,
    analysis_report,
    filter_events,
    filter_products,
    kpis,
    monthly,
    options,
    scenario_comparison,
)
from analyst.case_study.brief import number, priority_block
from analyst.case_study.config import (
    MASTER_PATH,
    REQUIRED_EVENT_FIELDS,
    REQUIRED_MASTER_FIELDS,
    SOURCE_URL,
)
from analyst.case_study.ingestion import fetch_events
from analyst.case_study.models import Snapshot
from analyst.case_study.preparation import build_snapshot
from analyst.case_study.workbook import generate_workbook
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
REPORT_SCHEMA = "canonical-submission-v1"
ASSIGNMENT_COMMIT = "f6adf2451ac6297c1b749abc131cff8a46020dc9"
ASSIGNMENT_RAW = "https://raw.githubusercontent.com/phongtdt/RFID_Data-Analyst_Case-Study-Oct-2026"
START_MARKER = "<!-- BEGIN GENERATED CANONICAL ANALYSIS -->"
END_MARKER = "<!-- END GENERATED CANONICAL ANALYSIS -->"
SCENARIOS = ("exclude_cost_10", "ambiguous_dm", "year_2024")
REVIEW_PRODUCT_IDS = frozenset(range(81, 86))


def sha256(payload: bytes) -> str:
    return hashlib.sha256(payload).hexdigest()


def json_bytes(value: Any) -> bytes:
    return (json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False) + "\n").encode("utf-8")


def code_identity() -> dict[str, str]:
    paths = sorted((ROOT / "analyst" / "case_study").glob("*.py")) + [Path(__file__)]
    return {
        path.relative_to(ROOT).as_posix(): sha256(path.read_bytes().replace(b"\r\n", b"\n"))
        for path in paths
    }


def remote_receipt(client: httpx.Client, source_path: str, expected_sha256: str) -> dict[str, Any]:
    url = f"{ASSIGNMENT_RAW}/{ASSIGNMENT_COMMIT}/{source_path}"
    response = client.get(url, headers={"Accept": "application/octet-stream"})
    response.raise_for_status()
    digest = sha256(response.content)
    if digest != expected_sha256:
        raise ValueError(
            f"Source differs from the verified assignment commit: {source_path}; independent re-audit required"
        )
    return {
        "url": url,
        "commit": ASSIGNMENT_COMMIT,
        "sha256": digest,
        "etag": response.headers.get("etag"),
        "bytes": len(response.content),
        "http_status": response.status_code,
        "fetched_at": datetime.now(UTC).isoformat(),
    }


def validate_real_snapshot(snapshot: Snapshot) -> None:
    provenance = snapshot.provenance
    if provenance.get("source_url") != SOURCE_URL or provenance.get("http_status") != 200:
        raise ValueError(
            "Submission generation requires a successful real assignment HTTP source; a fixture is not a source"
        )
    if not provenance.get("dataset_version") or not provenance.get("policy_version"):
        raise ValueError("Submission snapshot must carry dataset_version and policy_version")
    if provenance.get("scenario", ""):
        raise ValueError(
            "The submission requires the canonical policy, without a selected scenario"
        )
    datetime.fromisoformat(provenance["fetched_at"].replace("Z", "+00:00"))


def reviewed_cohort(snapshot: Snapshot, analysis: dict[str, Any]) -> dict[str, Any]:
    """Reuse engine metrics for the fixed products 81-85 drill-down cohort."""
    scope = Scope(**analysis["scope"])
    events = [
        event for event in filter_events(snapshot, scope) if event.product_id in REVIEW_PRODUCT_IDS
    ]
    products = tuple(
        product
        for product in filter_products(snapshot, scope)
        if product.product_id in REVIEW_PRODUCT_IDS
    )
    ranks = {row["product_id"]: row["rank"] for row in analysis["products"]["intensity_all"]}
    return {
        "selection_rule": "Fixed drill-down cohort: the first five Category 4 product IDs, which include the four highest reference-intensity products. Chosen by the analyst; not named in the case brief, not a top-five ranking and not a cleaning exclusion.",
        "product_ids": sorted(REVIEW_PRODUCT_IDS),
        "scope": analysis["scope"],
        "kpis": kpis(events, products, len(scope.months())),
        "monthly": monthly(events, products, scope, options(snapshot)),
        "products": [
            {**row, "reference_intensity_rank": ranks[row["product_id"]]}
            for row in analysis["products"]["all"]
            if row["product_id"] in REVIEW_PRODUCT_IDS
        ],
    }


def build_report_snapshot(snapshot: Snapshot, receipts: dict[str, Any]) -> dict[str, Any]:
    validate_real_snapshot(snapshot)
    analysis = analysis_report(snapshot, Scope())
    scenario_results = {}
    for identifier in SCENARIOS:
        result = scenario_comparison(snapshot, Scope(), identifier)
        selected = result["selected"]
        scenario_results[identifier] = {
            "scope": selected["scope"],
            "metadata": selected["metadata"],
            "kpis": selected["kpis"],
            "trend": selected["trend"],
            "comparison": result["comparison"],
            "recommendation_titles": [row["title"] for row in selected["recommendations"]],
            "reviewed_cohort": reviewed_cohort(snapshot, selected),
        }
    identities = code_identity()
    identity = {
        "report_schema_version": REPORT_SCHEMA,
        "dataset_version": snapshot.provenance["dataset_version"],
        "policy_version": snapshot.provenance["policy_version"],
        "scope": asdict(Scope()),
        "code_files": identities,
    }
    return {
        **identity,
        "report_id": sha256(json_bytes(identity))[:24],
        "generated_at": datetime.now(UTC).isoformat(),
        "source_receipts": receipts,
        "analysis": analysis,
        "raw_comparison": snapshot.raw_summary,
        "scenarios": scenario_results,
        "reviewed_cohort": reviewed_cohort(snapshot, analysis),
    }


def table(headers: list[str], rows: list[list[Any]]) -> list[str]:
    return ["| " + " | ".join(headers) + " |", "|" + "|".join("---" for _ in headers) + "|"] + [
        "| " + " | ".join(str(value).replace("|", "\\|").replace("\n", " ") for value in row) + " |"
        for row in rows
    ]


def render_analysis(report: dict[str, Any], language: str) -> str:
    """Render observed values from the shared snapshot; translations share paths."""
    vi = language == "vi"

    def n(value: int | float | None, decimals: int = 0) -> str:
        return number(value, language, decimals)

    def pct(value: float | None, decimals: int = 2) -> str:
        return "N/A" if value is None else n(value, decimals) + "%"

    def signed(value: float | None, decimals: int = 2) -> str:
        return "N/A" if value is None else ("+" if value > 0 else "") + n(value, decimals) + "%"

    def plus(value: int | float) -> str:
        return ("+" if value > 0 else "") + n(value)

    def pick(en: str, vn: str) -> str:
        return vn if vi else en

    a = report["analysis"]
    k, q, trend, halves = a["kpis"], a["quality"], a["trend"], a["period_halves"]
    location = a["location_severity"]
    totals = location["location_totals"]
    test = location["test"]
    largest = max(totals, key=lambda row: row["event_count"])
    severe = max(totals, key=lambda row: (row["critical_count"], row["event_count"]))
    line = next(row for row in a["lines"] if row["reference_intensity_per_1000"] is not None)
    volume_line = max(a["lines"], key=lambda row: row["event_count"])
    categories = [row for row in a["category_summary"] if row["product_count"] > 0]
    top_category = categories[0]
    others = categories[1:]
    control = next(
        (
            row
            for row in a["shared_category_controls"]
            if row["category"] == top_category["category"]
        ),
        None,
    )
    line_costs = a["line_cost_comparison"]
    lead_cost = line_costs[0]
    same_family = next(
        (
            row
            for row in (control or {}).get("lines", [])
            if row["production_line"] == lead_cost["production_line"]
        ),
        None,
    )
    recommendations = {row["kind"]: row for row in a["recommendations"]}
    category_gap = recommendations.get("category_intensity", {}).get("metrics", {})
    by_severity = {row["name"]: row for row in a["costs"]["by_severity"]}
    by_type = a["costs"]["by_type"]
    type_change = a["decomposition"]["defect_type"]
    line_change = a["decomposition"]["production_line"]
    top_type = type_change["rows"][0]
    top_line = next(row for row in line_change["rows"] if row["name"] != "Unknown")
    candidates = a["anomalies"]["candidates"]
    coverage = a["calendar_coverage"]
    issues = {issue["code"]: issue for issue in q["issues"]}
    dist = a["cost_distribution"]
    cluster = report["reviewed_cohort"]
    cohort_label = f"{min(cluster['product_ids'])}–{max(cluster['product_ids'])}"
    leading_products: list[int] = []
    for row in a["products"]["intensity_all"]:
        if row.get("category") != top_category["category"]:
            break
        leading_products.append(row["product_id"])
    first, second = halves["first"], halves["second"]
    month_span = (
        pick("January", "tháng 1") + "–" + pick("March", "tháng 3"),
        pick("April", "tháng 4") + "–" + pick("June", "tháng 6"),
    )
    severity_costs = ", ".join(
        f"{name} {n(by_severity[name]['average_known_repair_cost'], 2)}"
        for name in ("Minor", "Moderate", "Critical")
        if name in by_severity
    )
    other_events_per_product = [row["eligible_events_per_cohort_product"] for row in others]
    candidate_dates = ", ".join(row["date"] for row in candidates)
    single_line = [
        (row["category"], row["production_lines"][0])
        for row in categories
        if len(row["production_lines"]) == 1
    ]
    single_en = "; ".join(
        f"{category} is made only on {line_name}" for category, line_name in single_line
    )
    single_vi = "; ".join(
        f"{category} chỉ được sản xuất trên {line_name}" for category, line_name in single_line
    )
    content = [
        START_MARKER,
        "",
        pick("## Case study analysis: Part 3", "## Phân tích đề bài: Part 3"),
        "",
        pick(
            f"Scope: January–June 2024 under canonical policy `{report['policy_version']}`, dataset `{report['dataset_version']}` and report snapshot `{report['report_id']}`. The defect source was fetched over HTTP at `{a['metadata']['fetched_at']}`. Amounts are in **local currency (unspecified)**. Unrounded values, definitions and source receipts are in the [canonical snapshot](reports/canonical_snapshot.json); the dashboard shows the same numbers.",
            f"Phạm vi: tháng 1–6/2024 theo canonical policy `{report['policy_version']}`, dataset `{report['dataset_version']}` và snapshot báo cáo `{report['report_id']}`. Dữ liệu lỗi được lấy qua HTTP lúc `{a['metadata']['fetched_at']}`. Tiền tệ là **local currency (chưa xác định loại tiền)**. Số chưa làm tròn, định nghĩa và bằng chứng nguồn nằm trong [snapshot chuẩn](reports/canonical_snapshot.json); dashboard hiển thị đúng các số này.",
        ),
        "",
        pick("### Key findings", "### Phát hiện chính"),
        "",
        pick(
            f"1. **Defects are rising: a negative sign for the plant.** Logged events grew from {n(trend['first_count'])} in January to {n(trend['last_count'])} in June ({signed(trend['change_pct'])}). {month_span[1]} logged {n(second['event_count'])} events against {n(first['event_count'])} in {month_span[0]} ({signed(halves['event_count_change_pct'])}); Critical events rose {pct(halves['critical_count_change_pct'])} and observed repair cost {pct(halves['known_repair_cost_change_pct'])} between the two halves.",
            f"1. **Lỗi đang tăng: tín hiệu tiêu cực với nhà máy.** Số sự kiện lỗi tăng từ {n(trend['first_count'])} trong tháng 1 lên {n(trend['last_count'])} trong tháng 6 ({signed(trend['change_pct'])}). Giai đoạn {month_span[1]} ghi nhận {n(second['event_count'])} sự kiện so với {n(first['event_count'])} của {month_span[0]} ({signed(halves['event_count_change_pct'])}); giữa hai nửa giai đoạn, lỗi Critical tăng {pct(halves['critical_count_change_pct'])} và chi phí sửa chữa ghi nhận tăng {pct(halves['known_repair_cost_change_pct'])}.",
        ),
        pick(
            f"2. **{largest['location']} has the most defects; {severe['location']} has the most severe ones.** {largest['location']} logs {n(largest['event_count'])} events. {severe['location']} has {n(severe['critical_count'])} Critical events out of {n(severe['event_count'])} ({pct(severe['critical_share_pct'])}), against {pct(recommendations.get('critical_location', {}).get('metrics', {}).get('other_locations_critical_share_pct'))} at the other locations combined.",
            f"2. **{largest['location']} có nhiều lỗi nhất; {severe['location']} có nhiều lỗi nghiêm trọng nhất.** {largest['location']} có {n(largest['event_count'])} sự kiện. {severe['location']} có {n(severe['critical_count'])} lỗi Critical trên {n(severe['event_count'])} sự kiện ({pct(severe['critical_share_pct'])}), trong khi các vị trí còn lại gộp chung chỉ {pct(recommendations.get('critical_location', {}).get('metrics', {}).get('other_locations_critical_share_pct'))}.",
        ),
        pick(
            f"3. **{top_category['category']} has the highest observed reference intensity.** Its products log {n(top_category['reference_intensity_per_1000'], 3)} events per 1,000 reference units against {n(category_gap.get('rest_of_scope_reference_intensity_per_1000'), 3)} for all other products."
            + (
                f" Within each of the {n(control['lines_compared'])} lines that make them, they exceed the rest of the line by {n(control['min_intensity_ratio_to_rest_of_line'], 2)}–{n(control['max_intensity_ratio_to_rest_of_line'], 2)}×."
                if control
                else ""
            )
            + (
                f" The {n(len(leading_products))} highest-intensity products ({', '.join(str(item) for item in leading_products)}) are all {top_category['category']}."
                if len(leading_products) > 1
                else ""
            ),
            f"3. **{top_category['category']} có cường độ lỗi tham chiếu cao nhất.** Nhóm này có {n(top_category['reference_intensity_per_1000'], 3)} sự kiện trên 1.000 đơn vị sản lượng tham chiếu, so với {n(category_gap.get('rest_of_scope_reference_intensity_per_1000'), 3)} của toàn bộ sản phẩm còn lại."
            + (
                f" Trên cả {n(control['lines_compared'])} chuyền có sản xuất nhóm này, cường độ lỗi của nhóm gấp {n(control['min_intensity_ratio_to_rest_of_line'], 2)}–{n(control['max_intensity_ratio_to_rest_of_line'], 2)} lần phần còn lại của chuyền."
                if control
                else ""
            )
            + (
                f" {n(len(leading_products))} sản phẩm có cường độ cao nhất ({', '.join(str(item) for item in leading_products)}) đều thuộc {top_category['category']}."
                if len(leading_products) > 1
                else ""
            ),
        ),
        pick(
            f"4. **{lead_cost['production_line']} repairs cost {n(lead_cost['ratio_to_other_lines'], 2)}× more.** {lead_cost['production_line']} averages {n(lead_cost['average_known_repair_cost'], 2)} per known-cost repair against {n(lead_cost['other_lines_average_known_repair_cost'], 2)} on the other lines, so it carries {pct(lead_cost['known_cost_share_pct'])} of known repair cost with {pct(lead_cost['known_cost_event_share_pct'])} of the events.",
            f"4. **Sửa lỗi ở {lead_cost['production_line']} đắt gấp {n(lead_cost['ratio_to_other_lines'], 2)} lần.** Chi phí trung bình mỗi lần sửa có ghi nhận chi phí của {lead_cost['production_line']} là {n(lead_cost['average_known_repair_cost'], 2)}, so với {n(lead_cost['other_lines_average_known_repair_cost'], 2)} ở các chuyền khác; vì vậy chuyền này chiếm {pct(lead_cost['known_cost_share_pct'])} chi phí sửa chữa đã biết dù chỉ có {pct(lead_cost['known_cost_event_share_pct'])} số sự kiện.",
        ),
        pick(
            f"5. **The increase is concentrated.** Of the {plus(type_change['event_count_change'])} events between January and June, {top_type['name']} defects contribute {plus(top_type['event_count_change'])} ({n(top_type['contribution_pct_of_event_change'], 1)}%) and {top_line['name']} {plus(top_line['event_count_change'])} ({n(top_line['contribution_pct_of_event_change'], 1)}%).",
            f"5. **Mức tăng tập trung.** Trong {plus(type_change['event_count_change'])} sự kiện tăng thêm từ tháng 1 đến tháng 6, lỗi {top_type['name']} đóng góp {plus(top_type['event_count_change'])} ({n(top_type['contribution_pct_of_event_change'], 1)}%) và {top_line['name']} đóng góp {plus(top_line['event_count_change'])} ({n(top_line['contribution_pct_of_event_change'], 1)}%).",
        ),
        "",
        pick("### Data quality and counting units", "### Chất lượng dữ liệu và đơn vị đếm"),
        "",
        pick(
            f"- {n(q['raw_event_count'])} raw rows were received over HTTP. Cleaning removed {n(q['reconciliation']['duplicate_copies_removed'])} surplus copies ({n(q['exact_duplicate_count'])} exact, {n(q['normalized_duplicate_count'])} identical after label normalisation, {n(q['complementary_duplicate_count'])} missing-cost duplicate) and quarantined {n(q['quarantined_count'])} rows that share an ID but contradict each other, leaving {n(q['clean_event_count'])} events. {n(q['out_of_scope_count'])} valid events fall outside January–June 2024, so **{n(k['event_count'])} events are analysed**.",
            f"- Nhận {n(q['raw_event_count'])} dòng thô qua HTTP. Làm sạch loại {n(q['reconciliation']['duplicate_copies_removed'])} bản dư ({n(q['exact_duplicate_count'])} bản giống hệt, {n(q['normalized_duplicate_count'])} bản giống sau chuẩn hóa nhãn, {n(q['complementary_duplicate_count'])} bản trùng chỉ khác chi phí bị thiếu) và cách ly {n(q['quarantined_count'])} dòng cùng ID nhưng mâu thuẫn, còn {n(q['clean_event_count'])} sự kiện. {n(q['out_of_scope_count'])} sự kiện hợp lệ nằm ngoài tháng 1–6/2024, nên **phân tích dùng {n(k['event_count'])} sự kiện**.",
        ),
        pick(
            f"- {n(q['corrected_source_rows'])} rows were corrected ({n(issues.get('label_normalized', {}).get('affected_count'))} label case/whitespace fixes, {n(issues.get('day_first_date', {}).get('affected_count'))} unambiguous day-first dates). {n(q['raw_profile']['ambiguous_date_rows'])} slash dates could be read either way; they follow the dominant M/D/YYYY convention and a scenario tests the reverse. {n(q['flagged_in_scope_event_count'])} retained events carry at least one flag; flags and corrections overlap and are not additive.",
            f"- {n(q['corrected_source_rows'])} dòng được sửa ({n(issues.get('label_normalized', {}).get('affected_count'))} lỗi chữ hoa/khoảng trắng ở nhãn, {n(issues.get('day_first_date', {}).get('affected_count'))} ngày rõ ràng dạng ngày/tháng). {n(q['raw_profile']['ambiguous_date_rows'])} ngày dạng gạch chéo đọc được theo hai cách; policy theo quy ước M/D/YYYY chiếm đa số và có scenario thử cách đọc ngược lại. {n(q['flagged_in_scope_event_count'])} sự kiện được giữ có ít nhất một cờ; cờ và sửa đổi chồng lấp nên không cộng dồn.",
        ),
        pick(
            f"- {n(k['unmatched_event_count'])} events reference products missing from the master; they stay in plant totals but not in output-normalised rates. {n(k['cost_unknown_count'])} events have no valid cost (coverage {pct(k['cost_coverage_pct'])}); missing and negative costs are never imputed. {n(q['raw_profile']['cost_ten_rows'])} raw rows carry a cost of exactly 10, the lowest value observed; they stay valid and a scenario withholds them.",
            f"- {n(k['unmatched_event_count'])} sự kiện tham chiếu product không có trong master; chúng vẫn được tính vào tổng nhà máy nhưng không vào cường độ theo sản lượng. {n(k['cost_unknown_count'])} sự kiện không có chi phí hợp lệ (độ phủ {pct(k['cost_coverage_pct'])}); chi phí thiếu hoặc âm không bao giờ được điền ước tính. {n(q['raw_profile']['cost_ten_rows'])} dòng thô có chi phí đúng bằng 10, giá trị nhỏ nhất; chúng vẫn hợp lệ và có scenario tạm loại chúng.",
        ),
        pick(
            f"- No month has logged events after day {n(coverage['max_observed_day_of_month'])}, and {n(coverage['days_without_events'])} of {n(coverage['calendar_days'])} calendar dates have none. Reporting completeness is unverified, so month-end gaps must not be read as zero defects.",
            f"- Không tháng nào có sự kiện sau ngày {n(coverage['max_observed_day_of_month'])}, và {n(coverage['days_without_events'])}/{n(coverage['calendar_days'])} ngày lịch không có sự kiện nào. Mức đầy đủ của việc ghi nhận chưa được xác minh, nên không được hiểu các khoảng trống cuối tháng là không có lỗi.",
        ),
        pick(
            "- Every issue, its detection rule, the action taken and its impact are in [DATA_QUALITY.xlsx](DATA_QUALITY.xlsx).",
            "- Mọi vấn đề, cách phát hiện, cách xử lý và tác động nằm trong [DATA_QUALITY.xlsx](DATA_QUALITY.xlsx).",
        ),
        "",
        pick("### Q1. Trend, January–June 2024", "### Q1. Xu hướng tháng 1–6/2024"),
        "",
    ]
    content += table(
        [
            pick("Month", "Tháng"),
            pick("Defect events", "Sự kiện lỗi"),
            "Critical",
            pick("Observed repair cost", "Chi phí đã ghi nhận"),
            pick("Known cost / all events", "Có chi phí / tổng"),
            pick("Avg cost per known-cost event", "Chi phí TB mỗi sự kiện có chi phí"),
        ],
        [
            [
                row["month"],
                n(row["event_count"]),
                n(row["critical_count"]),
                n(row["known_repair_cost"], 2),
                f"{n(row['cost_known_count'])}/{n(row['event_count'])}",
                n(row["average_known_repair_cost"], 2),
            ]
            for row in a["monthly"]
        ],
    ) + [""]
    content += [
        pick(
            f"**Answer: a negative sign.** Logged events rose {pct(trend['change_pct'])} from January ({n(trend['first_count'])}) to June ({n(trend['last_count'])}). The later three months have a higher observed count: {month_span[1]} averaged {n(second['event_count'] / len(second['months']))} events a month against {n(first['event_count'] / len(first['months']))} in {month_span[0]} ({signed(halves['event_count_change_pct'])}). No month logs events after day {n(coverage['max_observed_day_of_month'])}, so reporting completeness must be checked before explaining February's dip or interpreting the rise as a change in the underlying process. Critical events ({signed(halves['critical_count_change_pct'])}) and observed repair cost ({signed(halves['known_repair_cost_change_pct'])}) rose with the count, and the average cost per known-cost repair climbed from {n(first['average_known_repair_cost'], 2)} to {n(second['average_known_repair_cost'], 2)}.",
            f"**Trả lời: tín hiệu tiêu cực.** Số sự kiện lỗi tăng {pct(trend['change_pct'])} từ tháng 1 ({n(trend['first_count'])}) đến tháng 6 ({n(trend['last_count'])}). Ba tháng cuối có số sự kiện cao hơn rõ rệt: {month_span[1]} trung bình {n(second['event_count'] / len(second['months']))} sự kiện mỗi tháng so với {n(first['event_count'] / len(first['months']))} của {month_span[0]} ({signed(halves['event_count_change_pct'])}). Không tháng nào có sự kiện sau ngày {n(coverage['max_observed_day_of_month'])}, nên cần kiểm tra mức đầy đủ của việc ghi nhận trước khi giải thích mức giảm tháng 2 hoặc coi mức tăng là thay đổi của quy trình sản xuất. Lỗi Critical ({signed(halves['critical_count_change_pct'])}) và chi phí sửa chữa ghi nhận ({signed(halves['known_repair_cost_change_pct'])}) tăng cùng số lỗi; chi phí trung bình mỗi lần sửa có chi phí tăng từ {n(first['average_known_repair_cost'], 2)} lên {n(second['average_known_repair_cost'], 2)}.",
        ),
        "",
        pick(
            f"Every line increased. {top_line['name']} grew fastest ({n(top_line['baseline']['event_count'])} → {n(top_line['comparison']['event_count'])} events) and contributes {n(top_line['contribution_pct_of_event_change'], 1)}% of the January–June increase; by defect type, {top_type['name']} contributes {n(top_type['contribution_pct_of_event_change'], 1)}% ({n(top_type['baseline']['event_count'])} → {n(top_type['comparison']['event_count'])}). A plant-wide rise points first to common causes such as volume, materials or inspection changes. Reference output is static, so output-normalised intensity rises in the same proportion; actual monthly output and inspected units are needed to confirm that the defect rate, not just production or inspection volume, went up.",
            f"Mọi chuyền đều tăng. {top_line['name']} tăng nhanh nhất ({n(top_line['baseline']['event_count'])} → {n(top_line['comparison']['event_count'])} sự kiện) và đóng góp {n(top_line['contribution_pct_of_event_change'], 1)}% mức tăng từ tháng 1 đến tháng 6; theo loại lỗi, {top_type['name']} đóng góp {n(top_type['contribution_pct_of_event_change'], 1)}% ({n(top_type['baseline']['event_count'])} → {n(top_type['comparison']['event_count'])}). Mức tăng trên toàn nhà máy gợi ý kiểm tra trước các nguyên nhân chung như sản lượng, vật liệu hoặc thay đổi kiểm tra. Sản lượng tham chiếu cố định nên cường độ theo sản lượng tăng cùng tỷ lệ; cần sản lượng thực tế và số đơn vị được kiểm theo tháng để khẳng định tỷ lệ lỗi tăng, chứ không chỉ sản lượng hay khối lượng kiểm tra tăng.",
        ),
        "",
        pick("### Q2. Defect location and severity", "### Q2. Vị trí lỗi và mức độ nghiêm trọng"),
        "",
    ]
    cells = {(cell["location"], cell["severity"]): cell["count"] for cell in location["cells"]}
    content += table(
        [
            pick("Location", "Vị trí"),
            pick("Events", "Sự kiện"),
            pick("Share of events", "Tỷ trọng sự kiện"),
            "Minor",
            "Moderate",
            "Critical",
            pick("Critical share at location", "Tỷ trọng Critical tại vị trí"),
        ],
        [
            [
                row["location"],
                n(row["event_count"]),
                pct(row.get("share_pct")),
                n(cells.get((row["location"], "Minor"), 0)),
                n(cells.get((row["location"], "Moderate"), 0)),
                n(row["critical_count"]),
                pct(row["critical_share_pct"]),
            ]
            for row in totals
        ],
    ) + [""]
    critical_metrics = recommendations.get("critical_location", {}).get("metrics", {})
    content += [
        pick(
            f"**Answer:** {largest['location']} is the most frequent location ({n(largest['event_count'])} events). Severity depends on location: {severe['location']}'s Critical share is {pct(severe['critical_share_pct'])}, against {pct(critical_metrics.get('other_locations_critical_share_pct'))} elsewhere. The association is statistically clear (chi-square p = {n(test['p_value'], 6)}; expected-cell checks {'pass' if test['reliable'] else 'fail'}) but small in size (Cramér's V = {n(test['cramers_v'], 3)}).",
            f"**Trả lời:** {largest['location']} là vị trí có nhiều lỗi nhất ({n(largest['event_count'])} sự kiện). Mức độ nghiêm trọng phụ thuộc vị trí: tỷ trọng Critical của {severe['location']} là {pct(severe['critical_share_pct'])}, so với {pct(critical_metrics.get('other_locations_critical_share_pct'))} ở các vị trí khác. Mối liên hệ có ý nghĩa thống kê rõ (chi-square p = {n(test['p_value'], 6)}; điều kiện tần suất kỳ vọng {'đạt' if test['reliable'] else 'không đạt'}) nhưng mức độ nhỏ (Cramér's V = {n(test['cramers_v'], 3)}).",
        ),
        "",
        pick(
            f"**Where to focus:** start failure analysis with {severe['location']} Critical defects, because {severe['location']} holds {pct(critical_metrics.get('share_of_all_critical_pct'))} of all Critical events. {largest['location']} is the volume issue with the lowest Critical share, so standard containment and inspection fit it better than deep failure analysis. The location label does not identify a mechanism; teardown of representative failures and process data are needed to test causes.",
            f"**Nên tập trung vào đâu:** bắt đầu phân tích hỏng hóc với lỗi Critical tại {severe['location']}, vì vị trí này chiếm {pct(critical_metrics.get('share_of_all_critical_pct'))} tổng số lỗi Critical. {largest['location']} là vấn đề về số lượng với tỷ trọng Critical thấp nhất, nên phù hợp với biện pháp khoanh vùng và kiểm tra tiêu chuẩn hơn là phân tích hỏng hóc sâu. Nhãn vị trí không chỉ ra cơ chế hỏng; cần tháo mẫu lỗi đại diện và dữ liệu công đoạn để kiểm chứng nguyên nhân.",
        ),
        "",
        pick("### Q3. Other notable patterns", "### Q3. Các mẫu hình đáng chú ý khác"),
        "",
    ]
    patterns = [
        [
            pick(
                f"{top_category['category']} has the highest defect intensity on every line",
                f"{top_category['category']} có cường độ lỗi cao nhất trên mọi chuyền",
            ),
            pick(
                f"{n(top_category['reference_intensity_per_1000'], 3)} vs {n(category_gap.get('rest_of_scope_reference_intensity_per_1000'), 3)} events per 1,000 reference units"
                + (
                    f"; {n(control['min_intensity_ratio_to_rest_of_line'], 2)}–{n(control['max_intensity_ratio_to_rest_of_line'], 2)}× the rest of each line"
                    if control
                    else ""
                )
                + f". Without the output denominator: {n(top_category['eligible_events_per_cohort_product'], 1)} events per product vs {n(min(other_events_per_product), 1)}–{n(max(other_events_per_product), 1)} for other categories. These are low-volume products ({n(top_category['average_monthly_output_per_product'])} units a month on average vs {n(k['known_monthly_output_subtotal'] / k['output_known_product_count'])} plant-wide).",
                f"{n(top_category['reference_intensity_per_1000'], 3)} so với {n(category_gap.get('rest_of_scope_reference_intensity_per_1000'), 3)} sự kiện trên 1.000 đơn vị tham chiếu"
                + (
                    f"; gấp {n(control['min_intensity_ratio_to_rest_of_line'], 2)}–{n(control['max_intensity_ratio_to_rest_of_line'], 2)} lần phần còn lại của từng chuyền"
                    if control
                    else ""
                )
                + f". Không dùng mẫu số sản lượng: {n(top_category['eligible_events_per_cohort_product'], 1)} sự kiện mỗi sản phẩm so với {n(min(other_events_per_product), 1)}–{n(max(other_events_per_product), 1)} ở các nhóm khác. Đây là sản phẩm sản lượng thấp (trung bình {n(top_category['average_monthly_output_per_product'])} đơn vị/tháng so với {n(k['known_monthly_output_subtotal'] / k['output_known_product_count'])} của toàn nhà máy).",
            ),
            pick(
                "The category has elevated reference intensity across the lines that make it. Shared design, materials and changeover practices are hypotheses to investigate; these descriptive comparisons do not separate product and line causes. Verify actual exposure and differences among products before choosing a corrective action.",
                "Nhóm sản phẩm có cường độ tham chiếu cao trên các chuyền sản xuất nhóm này. Thiết kế, vật liệu và cách chuyển đổi sản xuất chung là các giả thuyết cần điều tra; so sánh mô tả này chưa tách được nguyên nhân từ sản phẩm và chuyền. Cần xác minh sản lượng thực tế và khác biệt giữa các sản phẩm trước khi chọn biện pháp khắc phục.",
            ),
        ],
        [
            pick(
                f"{lead_cost['production_line']} repairs cost much more",
                f"Chi phí sửa ở {lead_cost['production_line']} cao hơn hẳn",
            ),
            pick(
                f"{n(lead_cost['average_known_repair_cost'], 2)} vs {n(lead_cost['other_lines_average_known_repair_cost'], 2)} per known-cost repair ({n(lead_cost['ratio_to_other_lines'], 2)}×)"
                + (
                    f"; for {top_category['category']} alone, {n(same_family['average_known_repair_cost'], 2)} on {lead_cost['production_line']} vs {n(same_family['same_category_other_lines_average_known_repair_cost'], 2)} on other lines"
                    if same_family
                    else ""
                )
                + ".",
                f"{n(lead_cost['average_known_repair_cost'], 2)} so với {n(lead_cost['other_lines_average_known_repair_cost'], 2)} mỗi lần sửa có chi phí (gấp {n(lead_cost['ratio_to_other_lines'], 2)} lần)"
                + (
                    f"; riêng {top_category['category']}: {n(same_family['average_known_repair_cost'], 2)} ở {lead_cost['production_line']} so với {n(same_family['same_category_other_lines_average_known_repair_cost'], 2)} ở các chuyền khác"
                    if same_family
                    else ""
                )
                + ".",
            ),
            pick(
                f"This line accounts for the largest share of observed repair cost ({pct(lead_cost['known_cost_share_pct'])}). Investigate repair routes and cost recording alongside product mix: comparison within a category still leaves differences in products, failure types and severity unresolved.",
                f"Chuyền này chiếm tỷ trọng chi phí sửa chữa ghi nhận lớn nhất ({pct(lead_cost['known_cost_share_pct'])}). Cần điều tra cách sửa và cách ghi nhận chi phí cùng với cơ cấu sản phẩm: so sánh trong một nhóm vẫn chưa xử lý khác biệt về sản phẩm, loại lỗi và mức độ nghiêm trọng.",
            ),
        ],
        [
            pick("Line rankings depend on the denominator", "Thứ hạng chuyền phụ thuộc vào mẫu số"),
            pick(
                f"{volume_line['production_line']} has the most events ({n(volume_line['event_count'])}); {line['production_line']} has the highest intensity ({n(line['reference_intensity_per_1000'], 4)} per 1,000 reference units). "
                + (single_en + "." if single_en else ""),
                f"{volume_line['production_line']} có nhiều sự kiện nhất ({n(volume_line['event_count'])}); {line['production_line']} có cường độ cao nhất ({n(line['reference_intensity_per_1000'], 4)} trên 1.000 đơn vị tham chiếu). "
                + (single_vi + "." if single_vi else ""),
            ),
            pick(
                "Raw counts describe logged workload. Where a category exists on only one line, these data cannot separate line and product effects. Use shared-category comparisons where available, check remaining product-mix differences and collect actual production and inspection exposure before assessing line performance.",
                "Số đếm thô mô tả khối lượng lỗi ghi nhận. Khi một nhóm sản phẩm chỉ có ở một chuyền, dữ liệu này không tách được tác động của chuyền và sản phẩm. Có thể so sánh nhóm xuất hiện trên nhiều chuyền, kiểm tra chênh lệch cơ cấu sản phẩm còn lại và thu thập sản lượng thực tế cùng phạm vi kiểm tra trước khi đánh giá hiệu suất chuyền.",
            ),
        ],
        [
            pick(
                "Repair cost does not track severity",
                "Chi phí sửa không đi theo mức độ nghiêm trọng",
            ),
            pick(
                f"Average cost per known-cost event: {severity_costs}. By defect type, {by_type[0]['name']} averages {n(by_type[0]['average_known_repair_cost'], 2)} against {n(min(row['average_known_repair_cost'] for row in by_type[1:]), 2)}–{n(max(row['average_known_repair_cost'] for row in by_type[1:]), 2)}. Known-cost median {n(dist['quantiles']['median'], 2)}, p90 {n(dist['quantiles']['p90'], 2)}.",
                f"Chi phí trung bình mỗi sự kiện có chi phí: {severity_costs}. Theo loại lỗi, {by_type[0]['name']} trung bình {n(by_type[0]['average_known_repair_cost'], 2)} so với {n(min(row['average_known_repair_cost'] for row in by_type[1:]), 2)}–{n(max(row['average_known_repair_cost'] for row in by_type[1:]), 2)}. Trung vị chi phí đã biết {n(dist['quantiles']['median'], 2)}, p90 {n(dist['quantiles']['p90'], 2)}.",
            ),
            pick(
                "Cost cannot stand in for severity: prioritise safety and function by Critical counts and use cost to target repair-process work. A flat cost profile also suggests checking with Finance how repair cost is recorded.",
                "Không thể dùng chi phí thay cho mức độ nghiêm trọng: ưu tiên an toàn và chức năng theo số lỗi Critical, còn chi phí dùng để chọn việc cải tiến quy trình sửa. Chi phí gần như đồng đều cũng là lý do để kiểm tra với Tài chính cách ghi nhận chi phí sửa chữa.",
            ),
        ],
        [
            pick(
                "A historical daily screen highlights unusual counts",
                "Sàng lọc theo ngày chỉ ra các ngày có số lỗi cao bất thường",
            ),
            pick(
                f"{n(len(candidates))} days exceed the January-based limit of {n(a['anomalies']['upper_reference_limit'], 2)} events a day (January mean {n(a['anomalies']['baseline_daily_mean'], 2)}): {candidate_dates}.",
                f"{n(len(candidates))} ngày vượt ngưỡng {n(a['anomalies']['upper_reference_limit'], 2)} sự kiện/ngày tính từ tháng 1 (trung bình tháng 1 là {n(a['anomalies']['baseline_daily_mean'], 2)}): {candidate_dates}.",
            ),
            pick(
                "The first historical breach is in April. This retrospective screen shows which dates to investigate; it is not a real-time alert or a confirmed process change. Operational alerts need a rolling baseline per line and category, reliable reporting and validation against confirmed incidents.",
                "Ngày vượt ngưỡng đầu tiên trong dữ liệu lịch sử rơi vào tháng 4. Sàng lọc hồi cứu chỉ ra các ngày cần điều tra; đây chưa phải cảnh báo thời gian thực hay bằng chứng xác nhận quy trình đã thay đổi. Cảnh báo vận hành cần baseline trượt theo từng chuyền và nhóm sản phẩm, dữ liệu ghi nhận đáng tin cậy và kiểm chứng với sự cố đã xác nhận.",
            ),
        ],
        [
            pick("Data capture has gaps", "Việc ghi nhận dữ liệu còn khoảng trống"),
            pick(
                f"No events after day {n(coverage['max_observed_day_of_month'])} in any month; {n(coverage['days_without_events'])}/{n(coverage['calendar_days'])} days empty; {n(k['cost_unknown_count'])} unknown costs; {n(q['raw_profile']['cost_ten_rows'])} costs of exactly 10.",
                f"Không có sự kiện sau ngày {n(coverage['max_observed_day_of_month'])} ở bất kỳ tháng nào; {n(coverage['days_without_events'])}/{n(coverage['calendar_days'])} ngày trống; {n(k['cost_unknown_count'])} chi phí chưa biết; {n(q['raw_profile']['cost_ten_rows'])} chi phí đúng bằng 10.",
            ),
            pick(
                "Fix reporting before automating alerts; otherwise month-end gaps look like improvement and cost totals stay incomplete.",
                "Cần sửa việc ghi nhận trước khi tự động hóa cảnh báo; nếu không, khoảng trống cuối tháng sẽ trông như cải thiện và tổng chi phí vẫn thiếu.",
            ),
        ],
    ]
    content += table(
        [
            pick("Pattern", "Mẫu hình"),
            pick("Evidence", "Bằng chứng"),
            pick("Why it matters to the Production team", "Vì sao quan trọng với team Production"),
        ],
        patterns,
    ) + [""]
    content += [
        pick(
            f"**Drill-down: products {min(cluster['product_ids'])}–{max(cluster['product_ids'])}.** These are the first five {top_category['category']} product IDs and include the four highest-intensity products. Together they log {n(cluster['kpis']['event_count'])} events, {n(cluster['kpis']['critical_count'])} Critical and {n(cluster['kpis']['known_repair_cost'], 2)} observed repair cost over {n(cluster['kpis']['reference_output'])} reference units. Monthly events: "
            + "; ".join(f"{row['month']}: {n(row['event_count'])}" for row in cluster["monthly"])
            + ".",
            f"**Đi sâu: sản phẩm {min(cluster['product_ids'])}–{max(cluster['product_ids'])}.** Đây là năm mã sản phẩm đầu tiên của {top_category['category']} và gồm bốn sản phẩm có cường độ cao nhất. Cả nhóm có {n(cluster['kpis']['event_count'])} sự kiện, {n(cluster['kpis']['critical_count'])} Critical và {n(cluster['kpis']['known_repair_cost'], 2)} chi phí sửa chữa ghi nhận trên {n(cluster['kpis']['reference_output'])} đơn vị tham chiếu. Sự kiện theo tháng: "
            + "; ".join(f"{row['month']}: {n(row['event_count'])}" for row in cluster["monthly"])
            + ".",
        ),
        "",
    ]
    content += table(
        [
            pick("ID / Product", "ID / Sản phẩm"),
            pick("Events", "Sự kiện"),
            "Critical",
            pick("Events / 1,000 reference units", "Sự kiện / 1.000 tham chiếu"),
            pick("Observed cost", "Chi phí ghi nhận"),
            pick("Known cost / all events", "Có chi phí / tổng"),
            pick("Intensity rank (of 100)", "Hạng cường độ (trên 100)"),
        ],
        [
            [
                f"{row['product_id']} / {row['product_name']}",
                n(row["event_count"]),
                n(row["critical_count"]),
                n(row["reference_intensity_per_1000"], 3),
                n(row["known_repair_cost"], 2),
                f"{n(row['cost_known_count'])}/{n(row['event_count'])}",
                n(row["reference_intensity_rank"]),
            ]
            for row in cluster["products"]
        ],
    ) + [""]
    content += [pick("### Sensitivity to assumptions", "### Độ nhạy với giả định"), ""]
    scenario_rows = [
        [
            "Canonical",
            n(k["event_count"]),
            n(k["known_repair_cost"], 2),
            f"{n(k['cost_known_count'])}/{n(k['event_count'])}",
            n(cluster["kpis"]["event_count"]),
            n(cluster["kpis"]["critical_count"]),
            n(cluster["kpis"]["known_repair_cost"], 2),
        ]
    ]
    for identifier, result in report["scenarios"].items():
        selected = result["kpis"]
        selected_cluster = result["reviewed_cohort"]["kpis"]
        scenario_rows.append(
            [
                identifier,
                n(selected["event_count"]),
                n(selected["known_repair_cost"], 2),
                f"{n(selected['cost_known_count'])}/{n(selected['event_count'])}",
                n(selected_cluster["event_count"]),
                n(selected_cluster["critical_count"]),
                n(selected_cluster["known_repair_cost"], 2),
            ]
        )
    content += table(
        [
            "Scenario",
            pick("Events", "Sự kiện"),
            pick("Observed cost", "Chi phí ghi nhận"),
            pick("Known cost / all events", "Có chi phí / tổng"),
            pick(f"Products {cohort_label} events", f"Sự kiện SP {cohort_label}"),
            pick(f"Products {cohort_label} Critical", f"Critical SP {cohort_label}"),
            pick(f"Products {cohort_label} cost", f"Chi phí SP {cohort_label}"),
        ],
        scenario_rows,
    ) + [""]
    canonical_titles = [row["title"] for row in a["recommendations"]]
    changes = [
        identifier
        for identifier, result in report["scenarios"].items()
        if result["comparison"]["priority_order_changed"]
        or result.get("recommendation_titles", canonical_titles) != canonical_titles
    ]
    content += [
        pick(
            "`exclude_cost_10` withholds costs equal to 10 without changing event counts; `ambiguous_dm` reads the order-ambiguous dates as D/M; `year_2024` moves out-of-period years into 2024. Each scenario is calculated separately and never modifies the canonical result. "
            + (
                "Leading signals change for: " + ", ".join(changes) + "."
                if changes
                else "The leading signals (most severe location, highest-intensity line and category, costliest line per repair, top cost type and top cost product) and therefore the three priorities do not change in any tested scenario."
            )
            + " Differences show sensitivity to assumptions, not confidence intervals or savings.",
            "`exclude_cost_10` tạm loại các chi phí bằng 10 mà không đổi số sự kiện; `ambiguous_dm` đọc các ngày mơ hồ theo D/M; `year_2024` chuyển năm ngoài kỳ sang 2024. Mỗi scenario được tính riêng và không sửa kết quả canonical. "
            + (
                "Các tín hiệu dẫn đầu thay đổi ở: " + ", ".join(changes) + "."
                if changes
                else "Các tín hiệu dẫn đầu (vị trí nghiêm trọng nhất, chuyền và nhóm sản phẩm có cường độ cao nhất, chuyền có chi phí sửa cao nhất, loại lỗi và sản phẩm tốn chi phí nhất), và vì vậy cả ba ưu tiên, không đổi trong mọi scenario đã thử."
            )
            + " Chênh lệch cho biết độ nhạy với giả định, không phải khoảng tin cậy hay khoản tiết kiệm.",
        ),
        "",
        pick("## Part 4: Recommendations", "## Part 4: Đề xuất"),
        "",
        pick("### Top three priorities", "### Ba ưu tiên hàng đầu"),
        "",
        pick(
            "These are the same priorities the dashboard shows under “Top priorities for this view” in its Overview; the backend derives them from the analysis above. Targets are provisional until actual output and inspection volumes are available.",
            "Đây cũng là ba ưu tiên dashboard hiển thị ở mục “Ưu tiên hàng đầu cho phạm vi này” trong phần Tổng quan; backend suy ra chúng từ phân tích ở trên. Các mục tiêu là tạm thời cho đến khi có sản lượng và khối lượng kiểm tra thực tế.",
        ),
        "",
    ]
    for item in a["recommendations"]:
        content += priority_block(item, a, language) + [""]
    content += [pick("### Next steps for the tool", "### Bước tiếp theo cho công cụ"), ""]
    content += [
        pick(
            "Used every day, a static dashboard still leaves the Production team to notice changes, explain them and follow up by hand. The table maps each limitation to what this tool already does and the next capability to add.",
            "Khi dùng hằng ngày, dashboard tĩnh vẫn bắt team Production tự phát hiện thay đổi, tự giải thích và tự theo dõi bằng tay. Bảng dưới đối chiếu từng hạn chế với những gì công cụ này đã làm và năng lực nên bổ sung tiếp.",
        ),
        "",
    ]
    content += table(
        [
            pick("Limitation of a static dashboard", "Hạn chế của dashboard tĩnh"),
            pick("What this tool already does", "Công cụ này đã làm được"),
            pick("Next capability", "Năng lực bổ sung tiếp theo"),
        ],
        [
            [
                pick(
                    "It waits to be opened; nobody is told when something changes.",
                    "Phải tự mở ra xem; không ai được báo khi có thay đổi.",
                ),
                pick(
                    f"The Trends and unusual days section marks historical days above the January-based limit ({n(len(candidates))} in this full-case period); filters recalculate the displayed scope. It does not send real-time alerts.",
                    f"Phần Xu hướng và ngày bất thường đánh dấu các ngày lịch sử vượt ngưỡng tính từ tháng 1 ({n(len(candidates))} ngày trong toàn bộ kỳ đề bài); bộ lọc tính lại phạm vi hiển thị. Công cụ chưa gửi cảnh báo thời gian thực.",
                ),
                pick(
                    "Scheduled checks per shift, line and category that push an alert with its evidence to the owner (Teams, Zalo or e-mail).",
                    "Kiểm tra theo lịch từng ca, chuyền và nhóm sản phẩm, tự đẩy cảnh báo kèm bằng chứng đến người phụ trách (Teams, Zalo hoặc e-mail).",
                ),
            ],
            [
                pick(
                    "The weekly report is still assembled by hand.",
                    "Báo cáo tuần vẫn phải tổng hợp thủ công.",
                ),
                pick(
                    "The Dashboard export menu downloads CSV for the active filters and the separate full-case DATA_QUALITY.xlsx workbook. Markdown brief and evidence ZIP endpoints are backend capabilities, not current Dashboard controls.",
                    "Menu xuất dữ liệu của Dashboard tải CSV theo bộ lọc đang áp dụng và workbook DATA_QUALITY.xlsx riêng cho toàn bộ đề bài. Backend có endpoint brief Markdown và ZIP bằng chứng; Dashboard hiện không có nút cho các đầu ra này.",
                ),
                pick(
                    "Send the brief on a schedule to each department, with an approval step and delivery receipts.",
                    "Gửi bản tin theo lịch tới từng phòng ban, có bước duyệt và xác nhận đã nhận.",
                ),
            ],
            [
                pick(
                    "Follow-up questions need an analyst.",
                    "Câu hỏi đào sâu vẫn cần người phân tích.",
                ),
                pick(
                    "The Agent tab investigates questions with the same backend metrics and checks every number against tool receipts. It is a prototype: it needs an OpenAI key and its live acceptance suite has not fully passed ([AUDIT.md](AUDIT.md)).",
                    "Tab Agent điều tra câu hỏi bằng đúng các metric của backend và đối chiếu mọi con số với receipt của công cụ. Đây là bản prototype: cần OpenAI key và bộ kiểm thử live chưa đạt toàn bộ ([AUDIT.md](AUDIT.md)).",
                ),
                pick(
                    "Make it fast and reliable enough for daily use, then let it draft the investigation behind each alert automatically for a person to review.",
                    "Làm Agent đủ nhanh và ổn định cho dùng hằng ngày, sau đó để nó tự soạn bản điều tra cho mỗi cảnh báo để người phụ trách duyệt.",
                ),
            ],
            [
                pick(
                    "An insight does not become an action.",
                    "Phát hiện không tự trở thành hành động.",
                ),
                pick(
                    "Recommendations describe proposed actions and owners. The current interface has no follow-up task manager; backend records alone do not establish that an action was assigned or completed.",
                    "Các khuyến nghị nêu hành động và vai trò phụ trách đề xuất. Giao diện hiện không có trình quản lý công việc theo dõi; bản ghi backend không chứng minh công việc đã được giao hay hoàn thành.",
                ),
                pick(
                    "Shared corrective-action tracking with due dates, escalation and before/after measurement on the same definitions.",
                    "Theo dõi hành động khắc phục dùng chung, có hạn chót, escalation và đo trước/sau trên cùng định nghĩa.",
                ),
            ],
            [
                pick(
                    "Data arrives late and incomplete.",
                    "Dữ liệu đến muộn và thiếu.",
                ),
                pick(
                    "HTTP fetch with manual refresh; stale or failed refreshes are labelled.",
                    "Lấy dữ liệu qua HTTP và làm mới thủ công; lần làm mới lỗi hoặc dữ liệu cũ đều được gắn nhãn.",
                ),
                pick(
                    "An MES/ERP feed with freshness monitoring, enforced month-end logging and cost capture.",
                    "Nguồn MES/ERP có giám sát độ mới, bắt buộc ghi nhận đến cuối tháng và đủ chi phí.",
                ),
            ],
            [
                pick(
                    "Rates rest on assumptions.",
                    "Tỷ lệ còn dựa trên giả định.",
                ),
                pick(
                    "Reference intensity uses static monthly_output. Backend scenario calculations and the comparison in this report test date and cost assumptions; the Dashboard has no what-if controls.",
                    "Cường độ tham chiếu dùng monthly_output cố định. Phép tính scenario ở backend và bảng so sánh trong báo cáo kiểm tra giả định về ngày và chi phí; Dashboard không có bộ điều khiển thử giả định.",
                ),
                pick(
                    "Actual output and inspected units by product, shift and lot, plus machine and process step, to turn intensities into defect rates and test causes.",
                    "Sản lượng thực và số đơn vị được kiểm theo sản phẩm, ca và lô, cùng máy và công đoạn, để chuyển cường độ thành tỷ lệ lỗi thật và kiểm chứng nguyên nhân.",
                ),
            ],
        ],
    )
    content += [
        "",
        pick(
            "Owners are proposed roles and targets are provisional; no assignment has been sent and no plant outcome has been measured.",
            "Người phụ trách là vai trò đề xuất và mục tiêu là tạm thời; chưa có việc nào được giao và chưa đo được kết quả thực tế tại nhà máy.",
        ),
        "",
        END_MARKER,
    ]
    return "\n".join(content)


def replace_generated(text: str, generated: str) -> str:
    if text.count(START_MARKER) != 1 or text.count(END_MARKER) != 1:
        raise ValueError("README must contain exactly one canonical-analysis marker pair")
    return re.sub(
        re.escape(START_MARKER) + r".*?" + re.escape(END_MARKER),
        lambda _: generated,
        text,
        flags=re.DOTALL,
    )


def generate(output_dir: Path) -> dict[str, Any]:
    raw, provenance = fetch_events(SOURCE_URL)
    master_bytes = MASTER_PATH.read_bytes()
    template = ROOT / "DATA_QUALITY_TEMPLATE.xlsx"
    with httpx.Client(timeout=20, follow_redirects=False) as client:
        receipts = {
            "defects": {
                **remote_receipt(
                    client, "mock-api/defect_records.json", provenance["source_sha256"]
                ),
                "runtime_url": SOURCE_URL,
                "runtime_fetched_at": provenance["fetched_at"],
                "schema": list(REQUIRED_EVENT_FIELDS),
            },
            "master": {
                **remote_receipt(client, "data/master_data.csv", sha256(master_bytes)),
                "schema": list(REQUIRED_MASTER_FIELDS),
            },
            "template": {
                **remote_receipt(
                    client, "DATA_QUALITY_TEMPLATE.xlsx", sha256(template.read_bytes())
                ),
                "schema": ["Summary", "Issues Found", "Assumptions", "Impact on Results"],
            },
        }
    snapshot = build_snapshot(raw, provenance, master_bytes)
    report = build_report_snapshot(snapshot, receipts)
    rendered_readmes = {}
    for language, name in (("en", "README.md"), ("vi", "README.vi.md")):
        source = (output_dir / name) if (output_dir / name).exists() else ROOT / name
        rendered_readmes[name] = replace_generated(
            source.read_text(encoding="utf-8"), render_analysis(report, language)
        )
    files = ["README.md", "README.vi.md", "DATA_QUALITY.xlsx", "reports/canonical_snapshot.json"]
    output_dir.mkdir(parents=True, exist_ok=True)
    with TemporaryDirectory(prefix="submission-", dir=output_dir.parent) as temporary:
        staging = Path(temporary)
        reports_dir = staging / "reports"
        reports_dir.mkdir()
        (reports_dir / "canonical_snapshot.json").write_bytes(json_bytes(report))
        generated = generate_workbook(
            snapshot, staging, template, analysis=report["analysis"], report_snapshot=report
        )
        for name, rendered in rendered_readmes.items():
            (staging / name).write_text(rendered, encoding="utf-8", newline="\n")
        manifest = {
            "report_schema_version": REPORT_SCHEMA,
            "report_id": report["report_id"],
            "dataset_version": report["dataset_version"],
            "policy_version": report["policy_version"],
            "files": {name: sha256((staging / name).read_bytes()) for name in files},
            "workbook_delivery": generated["workbook_delivery"],
        }
        (reports_dir / "submission_manifest.json").write_bytes(json_bytes(manifest))
        verify(staging)
        for name in [*files, "reports/submission_manifest.json"]:
            target = output_dir / name
            target.parent.mkdir(parents=True, exist_ok=True)
            (staging / name).replace(target)
    return verify(output_dir)


def verify(output_dir: Path) -> dict[str, Any]:
    report = json.loads(
        (output_dir / "reports" / "canonical_snapshot.json").read_text(encoding="utf-8")
    )
    manifest = json.loads(
        (output_dir / "reports" / "submission_manifest.json").read_text(encoding="utf-8")
    )
    if report["report_schema_version"] != REPORT_SCHEMA or report["analysis"]["scope"] != asdict(
        Scope()
    ):
        raise ValueError("Saved submission has a different report schema or canonical scope")
    if report["analysis"]["metadata"]["source_url"] != SOURCE_URL:
        raise ValueError("Saved submission is not sourced from the assignment HTTP endpoint")
    if code_identity() != report["code_files"]:
        raise ValueError(
            "Calculation/report code changed since generation; regenerate the canonical submission"
        )
    local_inputs = {"master": MASTER_PATH, "template": ROOT / "DATA_QUALITY_TEMPLATE.xlsx"}
    for name, path in local_inputs.items():
        if sha256(path.read_bytes()) != report["source_receipts"][name]["sha256"]:
            raise ValueError(f"Original {name} differs from the saved source receipt")
    for name, expected in manifest["files"].items():
        if sha256((output_dir / name).read_bytes()) != expected:
            raise ValueError(f"Saved deliverable hash differs from its manifest: {name}")
    for language, name in (("en", "README.md"), ("vi", "README.vi.md")):
        text = (output_dir / name).read_text(encoding="utf-8")
        if replace_generated(text, render_analysis(report, language)) != text:
            raise ValueError(f"README analysis differs from canonical snapshot: {name}")
    a, q = report["analysis"], report["analysis"]["quality"]
    workbook = load_workbook(output_dir / "DATA_QUALITY.xlsx", data_only=False)
    try:
        if workbook.sheetnames[:4] != [
            "Summary",
            "Issues Found",
            "Assumptions",
            "Impact on Results",
        ]:
            raise ValueError("Completed workbook lost supplied template sheets or order")
        expected_summary = [
            q["raw_event_count"],
            q["reconciliation"]["duplicate_copies_removed"] + q["quarantined_count"],
            q["corrected_source_rows"],
            q["flagged_in_scope_event_count"],
            q["in_scope_event_count"],
        ]
        if [workbook["Summary"].cell(row, 2).value for row in range(4, 9)] != expected_summary:
            raise ValueError("Completed workbook Summary differs from canonical snapshot")
        for row_number, row in enumerate(a["monthly"], start=2):
            observed = [
                workbook["Monthly Results"].cell(row_number, column).value for column in range(1, 8)
            ]
            expected = [
                row["month"],
                row["event_count"],
                row["critical_count"],
                row["critical_share_pct"],
                row["known_repair_cost"],
                row["cost_known_count"],
                row["cost_unknown_count"],
            ]
            if any(
                (
                    not isinstance(got, (int, float))
                    or not math.isclose(got, wanted, rel_tol=1e-14, abs_tol=1e-12)
                )
                if isinstance(wanted, float)
                else got != wanted
                for got, wanted in zip(observed, expected, strict=True)
            ):
                raise ValueError(f"Completed workbook monthly result differs: {row['month']}")
        provenance = {
            row[0].value: row[1].value for row in workbook["Provenance"].iter_rows(min_row=2)
        }
        if provenance.get("Canonical report ID") != report["report_id"]:
            raise ValueError("Completed workbook belongs to a different canonical report")
        action_rows = [
            row for row in workbook["Issues Found"].iter_rows(min_row=4) if row[2].value is not None
        ]
        rules = workbook["Issues Found"].data_validations.dataValidation
        if not all(
            row[5].value in {"Removed", "Corrected", "Flagged"}
            and any(
                rule.type == "list"
                and rule.formula1 == '"Removed,Corrected,Flagged"'
                and row[5].coordinate in rule
                for rule in rules
            )
            for row in action_rows
        ):
            raise ValueError(
                "Completed workbook issue actions violate the supplied list validation"
            )
    finally:
        workbook.close()
    return {
        "status": "Passed",
        "report_id": report["report_id"],
        "dataset_version": report["dataset_version"],
        "policy_version": report["policy_version"],
        "checks": [
            "local source integrity",
            "current code identity",
            "deliverable hashes",
            "bilingual generated analysis",
            "template sheet order",
            "summary counting units",
            "monthly results",
            "canonical report identity",
            "action validation",
        ],
        "limitation": "Saved-snapshot verification does not fetch current HTTP data, certify visual layout or evaluate the live Agent.",
    }


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path, default=ROOT)
    parser.add_argument("--check", action="store_true")
    arguments = parser.parse_args()
    result = (
        verify(arguments.output_dir.resolve())
        if arguments.check
        else generate(arguments.output_dir.resolve())
    )
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
